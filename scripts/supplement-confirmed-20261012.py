"""Prepare two verified additions; read source with openpyxl, preserve package on merge."""
import csv,json,copy,zipfile,re,hashlib,sys
from pathlib import Path
from datetime import datetime
import xml.etree.ElementTree as ET
import openpyxl
sys.stdout.reconfigure(encoding='utf-8')
ROOT=Path(__file__).resolve().parents[1]
SRC=ROOT/'全国机器人开放课题_个人申报台账_检索补充版_20261010.xlsx'
TMP=ROOT/'.sync/excel_20261012'
OUT=TMP/'全国机器人开放课题_个人申报台账_检索补充版_20261012.xlsx'
NS='http://schemas.openxmlformats.org/spreadsheetml/2006/main'
REL='http://schemas.openxmlformats.org/officeDocument/2006/relationships'
PKG='http://schemas.openxmlformats.org/package/2006/relationships'
ET.register_namespace('',NS);ET.register_namespace('r',REL)
COMPAT={'mc':'http://schemas.openxmlformats.org/markup-compatibility/2006','x14ac':'http://schemas.microsoft.com/office/spreadsheetml/2009/9/ac','xr':'http://schemas.microsoft.com/office/spreadsheetml/2014/revision','xr2':'http://schemas.microsoft.com/office/spreadsheetml/2015/revision2','xr3':'http://schemas.microsoft.com/office/spreadsheetml/2016/revision3'}
for prefix,uri in COMPAT.items():ET.register_namespace(prefix,uri)
def q(n):return '{'+NS+'}'+n
def readcsv(p):return list(csv.DictReader(p.open(encoding='utf-8-sig')))[0]
def serial(s):return (datetime.fromisoformat(s)-datetime(1899,12,30)).days
def prep():
 w=openpyxl.load_workbook(SRC);detail=w.worksheets[2];headers=[c.value for c in detail[1]]
 ids={r[0] for r in list(detail.values)[1:]};assert len(ids)==401
 assert not {'T030-001','T4-002-001'}&ids
 a=readcsv(ROOT/'02_检索与核验记录/未覆盖平台续查_20261011/新增正式项目_030.csv')
 b=readcsv(ROOT/'02_检索与核验记录/第四步交叉补漏_20261011/新增正式项目_002.csv')
 personal='个人资格、限项及固定合作者落实留后续逐项核对；受理中不等于本人条件已满足。'
 A={h:None for h in headers};A.update({
 '项目编号':a['项目编号'],'项目类别/证据性质':'正式开放课题公告','依托单位':'安徽中汇规划勘测设计研究院股份有限公司（依托）；铜陵学院（合作）','实验室/项目':a['平台名称'],
 '申报状态':'受理中','项目年度':'2026','截止日期':serial('2026-10-20'),'北京时间截止时点':'正文未载具体时刻；受理2026-10-11至2026-10-20',
 '研究主题':a['实际相关任务'],'资助金额及类别':a['金额及名额'],'申请资格':a['申请条件'],
 '合作及访问':'负责人须为铜陵学院校外人员；组内至少1名实验室固定成员。鼓励跨单位联合申请。','经费使用及外拨':'明示立项拨1.5万元、取得2027阶段成果后拨余款；对校外单位拨款路径及经费使用细则未载。',
 '研究周期':a['期限'],'成果数量及署名':'资助型提交报告并满足四选一：①中文EI或SCI论文1篇，第一作者单位为铜陵学院，作者单位同时列实验室；②固定成员为通讯作者的中文EI或SCI论文2篇，作者单位含铜陵学院及实验室；③授权发明专利1件，铜陵学院列前六专利权人；④指导固定成员立项相关横向项目≥30万元，铜陵学院或实验室承接/参与并有阶段成果。论文不得投已被或曾被预警期刊。\n自筹型提交报告并满足三选一：①固定成员通讯作者的中文EI或SCI论文1篇，排除预警期刊；②发明专利1件或实用新型2件，专利权人含铜陵学院；③相关横向项目≥10万元，铜陵学院或实验室承接/参与。\n成果标注实验室基金及课题编号。',
 '知识产权':'成果归研究者与实验室共同所有，结题归档。','材料及提交方式':'截止前将签字盖章申请书扫描版发送联系人邮箱；不下载申请书、不开展咨询。',
 '本人条件确认':personal,'缺失条款及联系事项':'截止具体时刻、校外经费拨款路径未载；现有结题条件来自既存官方证据，不追加补查。',
 '公告/线索来源URL':a['来源URL'],'官方指南/办法/附件URL':a['来源URL'],'核查基准（北京时间）':'2026-10-12',
 '经费与成果条款摘要':'资助型最高3万元；自筹型不拨款。结题条件详见O列，申请条件详见K/L列。','地域排序组':'安徽／铜陵','联系入口':'沈震；2022013@tlu.edu.cn'})
 B={h:None for h in headers};B.update({
 '项目编号':b['项目编号'],'项目类别/证据性质':'正式开放课题共享公告（专题02）','依托单位':b['依托单位'],'实验室/项目':b['平台名称'],
 '申报状态':'受理中','项目年度':'2026','截止日期':serial('2026-10-20'),'北京时间截止时点':'正文未载具体时刻',
 '研究主题':b['实际相关任务'],'资助金额及类别':b['明示金额'],'申请资格':b['关键申请条件'],
 '合作及访问':'单项只支持一家单位申报，不支持联合申报；公告未载必须配置固定合作者。','经费使用及外拨':'公开正文未载经费外拨口径；上限不是个人现金额度。',
 '研究周期':'不超过12个月（公告申报须知）','成果数量及署名':'已取得公告正文和专题任务图未载具体结题验收数量、成果署名及考核条件；扫码指南细则未取得。',
 '知识产权':'已取得公开资料未载。','材料及提交方式':'申请书Word及签章PDF、保密审查证明PDF，发送对应实验室联系人邮箱；不下载申请书。',
 '本人条件确认':personal,'缺失条款及联系事项':'截止具体时刻、经费外拨和结题细则未取得/未载，按字段分别保留；不据此推测。',
 '公告/线索来源URL':b['来源URL'],'官方指南/办法/附件URL':b['来源URL'],'核查基准（北京时间）':'2026-10-12',
 '经费与成果条款摘要':'专题02各方向研究经费不高于20万元；结题条件未载，详见O列。','地域排序组':'四川／成都','联系入口':'马老师；02887557341/15882408092；ma_song@139.com'})
 additions=[A,B]
 coverage=[['2026-10-12已核项目补入','空间信息获取与应用安徽省联合共建学科重点实验室','安徽／铜陵','2026开放课题：复用T030-001；10月12日官方正文复核；低空无人机多传感器集成、智能解译及识别算法','截止时刻及校外经费拨款路径未载','当期公告与具体任务已核，不等于机构全部平台覆盖',a['来源URL'],'仅保存具体缺口，不重复查既有批次'],
 ['2026-10-12已核项目补入','中电天奥·智能测控与天基信息应用四川省重点实验室','四川／成都','2026专题02共享公告：复用T4-002-001；窗口10月20日；无人机感知博弈、地面站自主及运动测量算法','截止时刻、外拨及结题细则未载/未取得','当期公告及既存任务图已核，不按方向数增加项目',b['来源URL'],'具体缺口保留，不推测结题要求']]
 payload={'source_sha256':hashlib.sha256(SRC.read_bytes()).hexdigest(),'headers':headers,'details':[[d[h] for h in headers] for d in additions],
 'overview':[['受理中',serial('2026-10-20'),'安徽中汇规划勘测设计研究院股份有限公司／铜陵学院·'+a['平台名称'],'低空无人机多传感器集成；遥感影像智能解译、多目标识别与自动监测（低空遥感场景）','资助型最高3万元/项；自筹型不拨款','公告'],
 ['受理中',serial('2026-10-20'),'中电天奥有限公司·'+b['平台名称'],'天基测控无人机感知博弈；地面站任务解析与流程自主；卫星前向散射运动参数测量（天基测控场景）','各方向研究经费≤20万元；外拨未载','公告']],
 'urls':[a['来源URL'],b['来源URL']],'coverage':coverage}
 TMP.mkdir(parents=True,exist_ok=True);(TMP/'additions.json').write_text(json.dumps(payload,ensure_ascii=False,indent=2),encoding='utf-8')
 record=ROOT/'02_检索与核验记录/第四步交叉补漏_20261011/main006_20261012/本次Excel补入输入.json'
 record.write_text(json.dumps({'核查日':'2026-10-12','用户最新范围':'机器人算法相关；来源、截止日期、具体方向、金额、申请条件、结题条件','补入':additions,'未补入':'氢储运加注装备公告缺明确机器人算法任务联系，留原证据；不作为新增相关项目'},ensure_ascii=False,indent=2),encoding='utf-8')
 print('prepared 2 rows; source 24 overview / 401 detail; no workbook edit')
def xmlwrite(e):
 ET.register_namespace('',PKG if e.tag.startswith('{'+PKG+'}') else NS)
 text=ET.tostring(e,encoding='unicode')
 if '{'+COMPAT['mc']+'}Ignorable' in e.attrib:
  end=text.index('>');first=text[:end]
  for prefix in e.attrib['{'+COMPAT['mc']+'}Ignorable'].split():
   if 'xmlns:'+prefix+'=' not in first:first+=' xmlns:'+prefix+'="'+COMPAT[prefix]+'"'
  text=first+text[end:]
 return ('<?xml version="1.0" encoding="UTF-8"?>\n'+text).encode('utf-8')
def merge():
 p=json.loads((TMP/'additions.json').read_text(encoding='utf-8'));assert hashlib.sha256(SRC.read_bytes()).hexdigest()==p['source_sha256']
 with zipfile.ZipFile(SRC) as z: parts={n:z.read(n) for n in z.namelist()}
 with zipfile.ZipFile(TMP/'delta.xlsx') as z:
  shared=[]
  if 'xl/sharedStrings.xml' in z.namelist():shared=[''.join(e.itertext()) for e in ET.fromstring(z.read('xl/sharedStrings.xml'))]
  ds=[ET.fromstring(z.read('xl/worksheets/sheet%d.xml'%i)) for i in [1,2,3]]
 def rows(e):return list(e.find(q('sheetData')))
 def rowcopy(row,n,template):
  r=copy.deepcopy(row);r.attrib.update(template.attrib);r.set('r',str(n));r.set('ht','75' if len(r)<=8 else '300');r.set('customHeight','1')
  for i,c in enumerate(r):
   c.set('r',re.sub(r'\d+$',str(n),c.get('r')))
   col=openpyxl.utils.cell.column_index_from_string(re.match(r'[A-Z]+',c.get('r')).group())
   tc=next((x for x in template if re.match(r'[A-Z]+',x.get('r')).group()==openpyxl.utils.get_column_letter(col)),None)
   c.set('s',tc.get('s','0') if tc is not None else '10')
   if c.get('t')=='s':
    v=c.find(q('v'));text=shared[int(v.text)];c.remove(v);c.set('t','inlineStr');is_=ET.SubElement(c,q('is'));ET.SubElement(is_,q('t')).text=text
  return r
 # Details append only; all original rows and source URLs remain.
 det=ET.fromstring(parts['xl/worksheets/sheet3.xml']);dat=det.find(q('sheetData'));template=dat[1]
 detail_numbers=[]
 for i,row in enumerate(rows(ds[1])):
  n=403+i;detail_numbers.append(n);dat.append(rowcopy(row,n,template))
 det.find(q('dimension')).set('ref','A1:AA404');parts['xl/worksheets/sheet3.xml']=xmlwrite(det)
 # Existing overview rows retain exact contents/styles; insert the two rows after N056.
 ov=ET.fromstring(parts['xl/worksheets/sheet1.xml']);data=ov.find(q('sheetData'));old=list(data);newrows=rows(ds[0]);mapping={};insert_at=12
 for row in old:data.remove(row)
 n=1;newrefs=[]
 for row in old:
  oldnum=int(row.get('r'));mapping[oldnum]=n
  row.set('r',str(n))
  for c in row:c.set('r',re.sub(r'\d+$',str(n),c.get('r')))
  data.append(row);n+=1
  if oldnum==insert_at:
   # Sensing/game and low-altitude recognition follow the existing perception/planning group.
   for idx in [1,0]:
    added=rowcopy(newrows[idx],n,old[1]);data.append(added);newrefs.append((idx,n));n+=1
 for h in ov.find(q('hyperlinks')):
  ref=h.get('ref');h.set('ref',re.sub(r'\d+$',str(mapping[int(re.search(r'\d+$',ref).group())]),ref))
 relpath='xl/worksheets/_rels/sheet1.xml.rels';rels=ET.fromstring(parts[relpath]);used={e.get('Id') for e in rels}
 hlinks=ov.find(q('hyperlinks'))
 for idx,n in newrefs:
  ET.SubElement(hlinks,q('hyperlink'),{'ref':f'C{n}','location':"'03所有项目完整详情'!A"+str(detail_numbers[idx]),'display':p['overview'][idx][2]})
  rid='rIdSupplement'+str(idx+1);assert rid not in used
  ET.SubElement(rels,'{'+PKG+'}Relationship',{'Id':rid,'Type':REL+'/hyperlink','Target':p['urls'][idx],'TargetMode':'External'})
  ET.SubElement(hlinks,q('hyperlink'),{'ref':f'F{n}','{'+REL+'}id':rid})
 ov.find(q('dimension')).set('ref','A1:F27');parts['xl/worksheets/sheet1.xml']=xmlwrite(ov);parts[relpath]=xmlwrite(rels)
 cov=ET.fromstring(parts['xl/worksheets/sheet4.xml']);cd=cov.find(q('sheetData'));ct=cd[6]
 for i,row in enumerate(rows(ds[2])):cd.append(rowcopy(row,2537+i,ct))
 cov.find(q('dimension')).set('ref','A1:H2538');parts['xl/worksheets/sheet4.xml']=xmlwrite(cov)
 # Only extend existing native filter tables; keep every header exactly as supplied.
 for name,ref in [('CurrentCalls','A1:F27'),('CompleteDetails','A1:AA404'),('InstitutionCoverage','A6:H2538')]:
  for path in list(parts):
   if path.startswith('xl/tables/') and path.endswith('.xml'):
    t=ET.fromstring(parts[path])
    if t.get('name')==name:
     t.set('ref',ref);a=t.find(q('autoFilter'))
     if a is not None:a.set('ref',ref)
     parts[path]=xmlwrite(t)
 with zipfile.ZipFile(OUT,'w',zipfile.ZIP_DEFLATED) as z:
  for k,v in parts.items():z.writestr(k,v)
 print('merged preserving all original row contents and header; output in local verification directory')
def verify():
 a=openpyxl.load_workbook(SRC);b=openpyxl.load_workbook(OUT)
 assert a.sheetnames==b.sheetnames
 def same(x,y):return x.value==y.value and x.style_id==y.style_id and x.number_format==y.number_format and str(x.hyperlink)==str(y.hyperlink)
 for k in [1,2,3]:
  for row in a.worksheets[k]:
   for c in row:assert same(c,b.worksheets[k][c.coordinate]),(k,c.coordinate)
 old=a.worksheets[0];new=b.worksheets[0]
 for row in old:
  n=row[0].row+(2 if row[0].row>12 else 0)
  for c in row:
   d=new.cell(n,c.column);assert c.value==d.value and c.style_id==d.style_id and c.number_format==d.number_format,(c.coordinate,d.coordinate)
   if c.hyperlink:assert c.hyperlink.target==d.hyperlink.target and c.hyperlink.location==d.hyperlink.location
 assert new['A1'].value==old['A1'].value
 assert len(list(b.worksheets[2].values))==404 and new.max_row==27
 assert [b.worksheets[2].cell(i,1).value for i in [403,404]]==['T030-001','T4-002-001']
 for s in b:
  assert s.freeze_panes==a[s.title].freeze_panes
  for t in s.tables.values():
   headerrow=int(re.search(r'\d+',t.ref).group());assert [s.cell(headerrow,i+1).value for i in range(len(t.tableColumns))]==[c.name for c in t.tableColumns]
 assert b.worksheets[2]['G403'].data_type=='d' and b.worksheets[2]['G404'].data_type=='d'
 report={'source_sha256':hashlib.sha256(SRC.read_bytes()).hexdigest(),'output_sha256':hashlib.sha256(OUT.read_bytes()).hexdigest(),'sheets':b.sheetnames,'overview_before':24,'overview_after':26,'details_before':401,'details_after':403,'existing_values_styles_links_preserved':True,'headers_unchanged':True,'table_header_consistency':True,'native_excel_verification':'pending'}
 (ROOT/'02_检索与核验记录/第四步交叉补漏_20261011/main006_20261012/Excel补入内容校验.json').write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf-8')
 print(json.dumps(report,ensure_ascii=False))
if __name__=='__main__':{'prep':prep,'merge':merge,'verify':verify}[sys.argv[1]]()
