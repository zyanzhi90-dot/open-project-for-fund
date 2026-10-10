from pathlib import Path
from openpyxl import load_workbook
from openpyxl.utils.datetime import to_excel
import re,json,hashlib,datetime
H=Path(__file__).parent;R=H.parents[2];source=R/'全国机器人开放课题_个人申报台账_兼容性修正版_20261010.xlsx'
w=load_workbook(source);details=list(w['02项目完整详情'].values);items={}
for sn in ['01申报总览','04其他资助']:
 s=w[sn]
 for n in range(2,s.max_row+1):
  rr=int(re.search(r'A(\d+)',s.cell(n,3).hyperlink.location)[1]);ident=details[rr-1][0];vs=[c.value for c in s[n]][:6]
  if isinstance(vs[1],datetime.datetime):vs[1]=to_excel(vs[1])
  vs[5]='公告';items[ident]={'id':ident,'source_sheet':sn,'source_row':n,'detail_row':rr,'url':s.cell(n,6).hyperlink.target,'values':vs}
order=['M078','M029','M066','M026','M077','M105','R001','N164','N056','N034','M160','M041','S001','M015','F002','M100','M161','M017','M142','F022','N085']
reasons={
 'M078':'医疗机器人力触觉反馈与人机协同，直接联系力觉感知及人机交互主线。',
 'M029':'机器人技能学习、自主控制和灵巧操作，联系机器人学习、规划与控制。',
 'M066':'多模态视觉融合、混杂系统建模与反馈镇定、决策协同控制，覆盖感知—建模—决策—控制方法；原公告未限定机器人应用。',
 'M026':'智能机电系统动力学建模、多源感知融合及先进控制，联系交互系统建模与反馈控制；保留电液系统背景。',
 'M077':'机器人灵巧操作与自主导航、视觉—语言—动作模型，联系感知、学习与运动规划。',
 'M105':'检修作业具身智能与人机协同，联系交互操作；保留检修应用约束。',
 'R001':'多模态环境感知、具身自主系统与协同控制，联系感知与决策控制；保留轨道交通背景。',
 'N164':'三维视觉定位、多模态感知、协同调度与路径规划；保留农业机器人和驻场要求。',
 'N056':'三维重建、定位导航和协同飞行，联系几何感知与运动规划；保留测绘和低空背景。',
 'N034':'受限环境感知、多源融合、救援路径规划与装备控制；保留矿山/隧道背景。',
 'M160':'装备智能检测与控制方法可联系系统反馈控制；原任务不明确力觉或柔顺交互。',
 'M041':'感知与推理决策方法有联系；脑启发任务未指定机器人交互系统。',
 'S001':'拓扑/几何规划及智能导航，联系路径规划；保留机载系统背景。',
 'M015':'安全约束三维航路规划，联系约束规划；原任务受通信质量和低空场景约束。',
 'F002':'机器人具身测控和边缘测控装备，联系感知测量；属于高校专项，原任务未明确接触力优化。',
 'M100':'运维机器人关键技术与系统集成；任务表述较宽，保留交通工程应用要求。',
 'M161':'机电液系统控制与机械动力学，联系建模控制；保留机械传动背景。',
 'M017':'六自由度测量补偿及水声通信，对应测量方法；与交互控制主线的直接联系较少。',
 'M142':'视觉检测及纺织具身智能有方法联系，但必须服务纺织研究；不按具身智能关键词与通用机器人任务等同排序。',
 'F022':'相关感知、无人装备和康复任务的合作参与渠道；青海牵头、转化应用、资金协议要求明确。',
 'N085':'可穿戴生物传感与健康监测涉及感知，但未列机器人操控/力觉/柔顺控制；保留服装应用原意。'}
assert set(order)==set(items) and len(order)==21
home=[items[i] for i in order]
removed=[]
for v in details[1:]:
 if str(v[0]).startswith('F') and v[0] not in ['F002','F022']:
  removed.append({'id':v[0],'name':v[3],'deadline':str(v[6]),'status':v[4],'destination':'03历史项目' if v[4]=='已截止' else '02项目完整详情','reason':v[18]})
d={'source':str(source),'source_sha256':hashlib.sha256(source.read_bytes()).hexdigest(),'output_name':'全国机器人开放课题_个人申报台账_申报合并与方向排序版_20261010.xlsx','home':home,'reasons':reasons,'removed_from_other':removed,'new_sheets':['01申报总览','02项目完整详情','03历史项目','04机构检索覆盖'],'proposal_basis_pages':[11,12,22,23],'detail_change':'none; user clarified no extra full-detail supplementation for F002/F022'}
(H/'plan.json').write_text(json.dumps(d,ensure_ascii=False,indent=2),encoding='utf8');print('Home merged 19+2=21; research-task order prepared; all existing details unchanged.');w.close()
