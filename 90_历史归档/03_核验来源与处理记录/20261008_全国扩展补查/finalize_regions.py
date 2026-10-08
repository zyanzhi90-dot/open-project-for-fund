"""Add native links/hidden states unsupported by the artifact authoring surface, then verify read-only."""
import pathlib,json,zipfile,xml.etree.ElementTree as E,hashlib,re,collections
from openpyxl import load_workbook
H=pathlib.Path(__file__).parent;ROOT=H.parents[2];D=json.loads((H/'prepared_regions.json').read_text(encoding='utf8'))
NS='http://schemas.openxmlformats.org/spreadsheetml/2006/main';REL='http://schemas.openxmlformats.org/officeDocument/2006/relationships';PKG='http://schemas.openxmlformats.org/package/2006/relationships'
E.register_namespace('',NS);E.register_namespace('r',REL)
tag=lambda s:'{'+NS+'}'+s
xml=lambda e:E.tostring(e,encoding='utf-8',xml_declaration=True)
with zipfile.ZipFile(H/'authored_regions.xlsx') as z:files={n:z.read(n) for n in z.namelist()}
book=E.fromstring(files['xl/workbook.xml']);sheetnums={s.attrib['name']:i+1 for i,s in enumerate(book.find(tag('sheets')))}
def links(name,seq):
 num=sheetnums[name];fn=f'xl/worksheets/sheet{num}.xml';rf=f'xl/worksheets/_rels/sheet{num}.xml.rels';s=E.fromstring(files[fn]);old=s.find(tag('hyperlinks'))
 if old is not None:s.remove(old)
 node=E.Element(tag('hyperlinks'));rels=E.fromstring(files[rf]) if rf in files else E.Element('{'+PKG+'}Relationships');ids={x.attrib['Id'] for x in rels};k=1000
 for cell,url,loc in seq:
  if not url and not loc:continue
  at={'ref':cell}
  if loc:at['location']=loc
  else:
   while 'rId'+str(k) in ids:k+=1
   rid='rId'+str(k);k+=1;ids.add(rid);at['{'+REL+'}id']=rid
   E.SubElement(rels,'{'+PKG+'}Relationship',{'Id':rid,'Type':REL+'/hyperlink','Target':url,'TargetMode':'External'})
  E.SubElement(node,tag('hyperlink'),at)
 later={'printOptions','pageMargins','pageSetup','headerFooter','rowBreaks','colBreaks','customProperties','cellWatches','ignoredErrors','smartTags','drawing','legacyDrawing','legacyDrawingHF','picture','oleObjects','controls','webPublishItems','tableParts','extLst'}
 idx=next((i for i,c in enumerate(s) if c.tag.split('}')[-1] in later),len(s));s.insert(idx,node);files[fn]=xml(s);files[rf]=xml(rels)
pm={p['id']:p for p in D['projects']};detailmap={p['id']:i+2 for i,p in enumerate(D['projects'])};logmap={}
for i,r in enumerate(D['logs']):logmap.setdefault(r[0],i+2)
for name,rows in D['list_maps'].items():
 seq=[]
 for x in rows:
  p=pm[x['id']];r=x['row'];seq.extend([(f'C{r}',None,f"'02项目完整详情'!A{detailmap[p['id']]}"),(f'F{r}',p['url'],None)])
 links(name,seq)
seq=[]
for i,p in enumerate(D['projects']):
 row=i+2;seq.extend([(f'A{row}',None,f"'04本轮信息核验日志'!A{logmap[p['id']]}"),(f'T{row}',p['url'],None)])
 u=next((u for u in p['sources'] if re.search(r'Download|\.pdf|\.doc|_upload/article/files',u,re.I)),None)
 if u:seq.append((f'U{row}',u,None))
links('02项目完整详情',seq)
links('04本轮信息核验日志',[(f'D{i+2}',r[3].splitlines()[0],None) for i,r in enumerate(D['logs']) if r[3].startswith('http')])
links('05剔除记录',[(f'F{i+2}',r[5],None) for i,r in enumerate(D['excluded']) if str(r[5]).startswith('http')])
for s in book.find(tag('sheets')):
 if s.attrib['name'].startswith(('90','91')):s.set('state','hidden')
files['xl/workbook.xml']=xml(book)
out=H/'全国机器人开放课题_个人申报台账_全国扩展版_20261009.xlsx'
with zipfile.ZipFile(out,'w',compression=zipfile.ZIP_DEFLATED) as z:
 for n,b in files.items():z.writestr(n,b)

assert hashlib.sha256(pathlib.Path(D['current_input']).read_bytes()).hexdigest()==D['current_input_hash']
w=load_workbook(out,data_only=False);B=json.loads((H/'base_rows.json').read_text(encoding='utf8'))
assert [c.value for c in w['01申报总览'][1]]==['申报状态','截止日期','依托单位·实验室','公告研究方向','资助金额','公告']
assert w['01申报总览'].max_column==6 and w['01申报总览'].max_row==D['stats']['current']+1
assert w['01申报总览'].freeze_panes=='A2' and w['02项目完整详情'].freeze_panes=='E2'
assert w['02项目完整详情'].max_row==D['stats']['total']+1
listed=[]
for name,rows in D['list_maps'].items():
 s=w[name];assert len(s.tables)==1 and s.max_row==len(rows)+1
 for x in rows:
  p=pm[x['id']];r=x['row'];listed.append(p['id']);assert s.cell(r,1).value==p['status']
  assert s.cell(r,3).hyperlink.location==f"'02项目完整详情'!A{detailmap[p['id']]}";assert s.cell(r,6).hyperlink.target==p['url']
  if p['date']:assert s.cell(r,2).value.strftime('%Y-%m-%d')==p['date']
  if name=='01申报总览':assert p['date']>='2026-10-09' and p['kind']=='正式开放课题公告'
assert len(set(listed))==len(listed)==D['stats']['total']
norm=lambda v: '' if v is None else (v.strftime('%Y-%m-%d') if hasattr(v,'strftime') else (v[:10] if isinstance(v,str) and 'T00:00:00' in v else v))
for name in ['90筛选前原始记录','91输入版本留档']:
 s=w[name];assert s.sheet_state=='hidden' and s.max_row==len(B[name])
 for i,row in enumerate(B[name],1):
  for j,v in enumerate(row,1):assert norm(s.cell(i,j).value)==norm(v),(name,i,j)
prior=load_workbook(D['current_input'],read_only=True,data_only=False)
changed=set(D['national_changes']['updated_ids'])
for row in prior['02项目完整详情'].iter_rows(min_row=2,values_only=True):
 if row[0] in changed:continue
 r=detailmap[row[0]]
 for j,v in enumerate(row,1):assert norm(w['02项目完整详情'].cell(r,j).value)==norm(v),('unchanged prior detail',row[0],j)
assert set(detailmap).issuperset(row[0] for row in B['02项目完整详情'][1:])
assert len(D['excluded'])==D['stats']['excluded']
for p in D['projects']:
 r=detailmap[p['id']]
 for j,v in enumerate(p['details'],1):
  actual=w['02项目完整详情'].cell(r,j).value
  if j==7 and not v:assert actual=='未取得'
  else:assert norm(actual)==norm(v),(p['id'],j)
for g in D['groups']:
 if g is D['groups'][1]:
  for id in g['ids']:
   if pm[id]['date']:assert pm[id]['date']<'2026-10-09',(id,pm[id]['date'])
assert [c.value for c in w['03机构检索覆盖'][6]]==B['03机构检索覆盖'][5]
assert w['04本轮信息核验日志'].max_row==len(D['logs'])+1
prior.close()
for s in w:
 for row in s:
  for c in row:assert c.data_type!='e',(s.title,c.coordinate)
for group in D['groups']:
 ranks=['P0 南京及周边','P0 澳门','P0 香港','P0 广州及周边','P1 北京及周边','P1 西安及周边','P1 重庆']
 keys=[(ranks.index(pm[i]['region']) if pm[i]['region'] in ranks else 7,pm[i]['date'] or '9999-99-99',pm[i]['org'],i) for i in group['ids']];assert keys==sorted(keys)
report={'output':str(out),'stats':D['stats'],'input_sha256_unchanged':True,'input_150_preserved_except_seven_documented_updates':True,'original_199_and_hidden_archives_unchanged':True,'home_six_columns_and_15_formal_open_calls':True,'unique_166_ids_dates_links_filters_freeze_and_order_valid':True,'sha256':hashlib.sha256(out.read_bytes()).hexdigest(),'visual_review':'pending'}
(H/'validation_regions.json').write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf8')
print(json.dumps(report,ensure_ascii=False))
