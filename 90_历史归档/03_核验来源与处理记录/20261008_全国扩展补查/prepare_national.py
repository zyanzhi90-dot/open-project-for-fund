import pathlib,json,copy,hashlib,re,datetime,collections,zipfile,xml.etree.ElementTree as E
from openpyxl import load_workbook
H=pathlib.Path(__file__).resolve().parent; ROOT=H.parents[2]; PREV=H.parent/'20261008_重点地区补查'
read=lambda p:json.loads(p.read_text(encoding='utf8'))
D=read(PREV/'prepared_regions.json'); B=read(H/'base_rows.json')
source=ROOT/'全国机器人开放课题_个人申报台账_重点地区补查版_20261008-网页修改版.xlsx'
w=load_workbook(source,data_only=False,read_only=True); D['current_input']=str(source); D['current_input_hash']=hashlib.sha256(source.read_bytes()).hexdigest()
old={p['id']:p for p in D['projects']}; projects=[]
for row in w['02项目完整详情'].iter_rows(min_row=2,values_only=True):
 r=[v.strftime('%Y-%m-%d') if isinstance(v,datetime.datetime) else v or '' for v in row]; p=copy.deepcopy(old.get(r[0],{}));
 p.update(dict(id=r[0],kind=r[1],org=r[2],lab=r[3],status=r[4],year=r[5],date=r[6],time=r[7],topic=r[8],amount=r[9],fields=r[10:17],personal=r[17],gaps=r[18],url=r[19],sources=list(dict.fromkeys([r[19]]+re.findall(r'https?://[^\s]+',r[20]))),relation=r[22],value=r[23],region=r[24],contact=r[25],details=r,aliases=p.get('aliases',[]),amountlong=r[9])); projects.append(p)
for p in projects:
 if not re.fullmatch(r'20\d\d-\d\d-\d\d',str(p['date'])):p['date']=''
D['projects']=projects; pm={p['id']:p for p in projects}; detailmap={i+2:p['id'] for i,p in enumerate(projects)}
NS={'m':'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
with zipfile.ZipFile(source) as z:
 for g,name in zip(D['groups'],['01申报总览','06历史项目','07待核项目','08其他资助']):
  num=w.sheetnames.index(name)+1; xml=E.fromstring(z.read(f'xl/worksheets/sheet{num}.xml'))
  loc={x.attrib['ref']:x.attrib.get('location','') for x in xml.findall('m:hyperlinks/m:hyperlink',NS)}
  g['ids']=[detailmap[int(re.search(r'!A(\d+)',loc[f'C{r}']).group(1))] for r in range(2,len(B[name])+1)]
 assert len(set(id for g in D['groups'] for id in g['ids']))==150
for k,s in [('coverage','03机构检索覆盖'),('archive_rows','90筛选前原始记录'),('prior_rows','91输入版本留档')]:D[k]=copy.deepcopy(B[s])
D['logs']=copy.deepcopy(B['04本轮信息核验日志'][1:]);D['excluded']=copy.deepcopy(B['05剔除记录'][1:])
unknown='官方未公开（已取得正文未载）；管理附件未取得，不能推填。'
labels=['申请资格','合作及访问','经费使用及外拨','研究周期','成果数量及署名','知识产权','材料及提交方式']
changes=[]; new=[]
def rebuild(p):
 p['details']=[p['id'],p['kind'],p['org'],p['lab'],p['status'],p['year'],p['date'],p['time'],p['topic'],p['amount'],*p['fields'],p['personal'],p['gaps'],p['url'],'\n'.join(p['sources']),'2026-10-09（北京时间；检索执行于10月8—9日）',p['relation'],p['value'],p['region'],p['contact'],','.join(p['aliases'])];assert len(p['details'])==27
def log(p):
 for k,v in zip(labels,p['fields']):D['logs'].append([p['id'],k,v,'\n'.join(p['sources']),'2026-10-09','官方正文/附件；具体缺口逐字段标示',p.get('screen_decision','保留'),p['status']])
 for k,v in [('年度与受理',str(p['year'])+'；截止'+(p['date'] or '未取得')+' '+p['time']),('任务筛选',p['topic']),('金额',p['amount']),('手动补充事项',p['gaps'])]:D['logs'].append([p['id'],k,v,p['url'],'2026-10-09',p['kind'],p.get('screen_decision','保留'),p['status']])
def add(id,org,lab,region,year,date,topic,amount,url,group,fields,gaps,contact='',time='',sources=[]):
 assert id not in pm
 p=dict(id=id,org=org,lab=lab,region=region,year=year,date=date,topic=topic,amount=amount,amountlong=amount,url=url,kind='正式开放课题公告' if group<2 else '官方线索/长期机制',status=['受理中·条件待核','已截止','线索待核','其他资助'][group],fields=fields,gaps=gaps,contact=contact,time=time,sources=list(dict.fromkeys([url]+sources)),personal='已知：南京信息工程大学现职、博士。其余个人条件须按该项目原文确认；尚未落实固定合作者。',relation='',value='经费、成果和权属分别见M/O/P列；缺口见S列。',aliases=[],screen_decision='保留' if group<2 else '待核',screen_reason=topic)
 rebuild(p);D['projects'].append(p);pm[id]=p;D['groups'][group]['ids'].append(id);new.append(id);log(p)
def update(id,group=None,**values):
 p=pm[id];p.update(values);rebuild(p);log(p);changes.append(id)
 if group is not None:
  for g in D['groups']:
   if id in g['ids']:g['ids'].remove(id)
  D['groups'][group]['ids'].append(id)

add('N001','清华大学','北京信息科学与技术国家研究中心2026年度开放课题','P1 北京及周边','2026','2026-05-11','方向6：类脑物理智能体跨场景感知、认知、决策、执行；小样本迁移适配及平台验证','10万元/项；每方向1项','https://www.bnrist.tsinghua.edu.cn/info/1110/4271.htm',1,[
 '博士且中级以上职称；研究中心以外正式在编科研人员；博士后、学生不能申请；负责人及主要成员限申报1项。',
 '必须联合研究中心相关方向研究人员申请；保证在研究中心开展必要研究，天数未载。',
 '10万元；用于成员在中心开展相关科研；负责人及合作人共同签字管理。开支科目、管理费及能否外拨正文和已取得指南未载。',
 '1年：2026年6月—2027年6月。',
 '方向6要求形成验证平台；典型任务成功率提升5%以上，稀缺样本性能提升10%以上。方向6论文统一数量、单位署名与基金致谢未载；不得套方向5论文要求。',unknown,
 '纸质1份及Word电子版；电子名2026开放课题申请-群体名称-申请人姓名；xuge@tsinghua.edu.cn，通知另列wuky@tsinghua.edu.cn。纸质FIT楼3-322，100084。截止5月11日，时点未载。'
 ],'经费科目/管理费/外拨、访问天数、方向6署名与知识产权未载，须咨询；已取得年度PDF，不需再找指南。','徐老师010-62795788；吴老师010-62797486；xuge@tsinghua.edu.cn',sources=['https://www.bnrist.tsinghua.edu.cn/__local/8/A8/DA/03EFA96344B085C4991D1A776E7_94C301CC_38E0B.pdf','https://www.bnrist.tsinghua.edu.cn/info/1110/4418.htm'])
add('N002','武汉软件工程职业学院／武汉开放大学','复杂零部件智能检测与识别湖北省工程研究中心','P2 湖北·武汉','2026','2026-06-30','工业机器人五轴扫描路径规划；AMR与协作臂、三维扫描闭环联动；自适应支脚力反馈调平','原则5万元/项；可按评审增加，上限未载','https://gcyjzx.whvcse.edu.cn/info/1041/3921.htm',1,[
 '国内外高校、研究机构及公司科研人员；博士或中级以上职称；负责人限1项，在研未结题不得申请；参与限2项。','保证研究时间，具体天数未载；固定合作人员硬性要求未载。',
 '依学校纵向科研项目管理办法及合同执行；科目、管理费、外拨与报销细则未取得。','一般不超过2年。',
 '至少满足七类成果中的两类：关键技术、新产品及收益、科技奖、标准、省部级项目/横向到账、发明专利或软著、全国中文核心以上论文。具体单类数量未载，不能改成固定2篇论文。成果署中心；鉴定由中心组织，奖励共同申报。',
 '中心与负责人所在单位共享。','电子申请书及汇总表2026-06-30 24:00前至gcyjzx2025@163.com；签章纸质材料7月10日前报送，具体份数及收件位置需完整附件。'
 ],'官方正文取得主要条款；附件和正文末尾具体纸质份数/收件地址未取得。另一个学校官方转载域名亦访问失败，停止尝试。','gcyjzx2025@163.com',time='24:00（北京时间）',sources=['https://www.whvcse.edu.cn/info/2289/209518.htm'])
add('N003','同济大学','国家土建结构预制装配化工程技术研究中心沈祖炎专项开放基金','P2 上海','2026','2026-06-30','定向方向8：村镇模块化钢结构体系，以及焊接机器人适配工艺、路径和参数；同时包括建筑被动式降温任务','机器人相关方向为定向项目20万元；非定向每项7.5万元另列，不混填','https://prefabcenter.tongji.edu.cn/c0/f8/c15384a377080/page.htm',1,[
 '定向项目优先合作机构；非合作单位须通过中心协调，事先与资助方达成合作意向。非定向另列博士/副高、年龄≤40等，是否适用于机器人定向方向8须确认，不能直接套用。',
 '至少1名中心固定成员参与；依托单位同意；未载具体访问天数。',unknown,'2年：2026-09-01—2028-08-31。',
 '定向考核由中心和资助方协商；不套非定向的论文/专利数量。中期检查，结束60天内交报告与决算；延期提前90天申请，最长1年。',unknown,
 '签章纸质4份及一致PDF发liyan0714@tongji.edu.cn；获批任务书4份；上海四平路1239号衷和楼2014，200092。'
 ],'机器人方向为定向项目，须确认接洽意向和其资格适用范围；经费外拨、权属、定向成果指标待管理办法/协议。','李艳；13808206416；liyan0714@tongji.edu.cn')
add('N004','浙江大学','全省农业智能感知与机器人重点实验室','P2 浙江·杭州','2025','2025-03-31','农业信息感知；作物收获与移栽装备；动物饲喂和产后处理机器人','2—3万元/项左右','https://www.nercita.org.cn/infos/detail?id=10341&type=203',1,[
 '中级及以上职称的国内外科研、教学与技术人员；鼓励来室合作。','获批执行期间为客座研究人员；固定人员联合与明确到访天数未载。',
 '专款专用，报销制、不外拨；本室研究实验费、到访差旅、相关资料与学术活动费。','2年，从批准日起执行。',
 '相关SCI/EI收录论文1篇；每年进展，期满3个月内结题；论文须含实验室单位署名。','实验室与申请人所在单位共享。',
 '申请书电子版至0921a09@zju.edu.cn；截止2025-03-31；纸质要求未载。'
 ],'取得北京市农林科学院官方转载全文；原实验室申请表未取得；2026—2027正式年度公告未检出，不沿用2025窗口。','屠雯雯0571-88982348；0921a09@zju.edu.cn')
add('N005','中北大学','山西省恶劣环境智能装备技术重点实验室','P2 山西·太原','2023','2023-07-30','开放方向（六）：机器人理论及技术；不以其他材料类方向作为保留依据','拟资助8—10项；单项金额未载','https://www.nuc.edu.cn/info/1014/28267.htm',1,[
 '国内外中级以上或博士科研人员；正文另列校内副高/博士中级和专家推荐途径，保留不同对象条件；外部合作/访问等科研人员可申请。','固定成员合作优先，并非正文一律强制；负责人每年限1项。',
 '合同后50%、中期检查后50%；开支科目、外拨及管理费未载，需办法。','2年，从2023年9月起执行。',
 '结题2篇高水平论文，或1篇高水平论文加1件发明专利申请；年度进展与次年计划、中期报告、答辩；按实验室署名要求。','已取得正文未载完整权属。',
 '2023-07-20—07-30；签字盖章纸质3份及电子1份；czy1393518@163.com；太原学院路3号深孔楼2层。'
 ],'2023年度金额、外拨及权属管理办法未取得；2026—2027本平台指南未检出。','陈振亚13935181410；czy1393518@163.com')
add('N006','阳光学院','福建省空间信息感知与智能处理重点实验室','P2 福建·福州','2022','2022-11-04','机器人导航：SLAM与视觉导航、雷达视觉融合、深度学习路径规划、ROS移动机器人开发','一般2万元；重点4—6万元','https://isdm.ygu.edu.cn/info/1100/1376.htm',1,[
 '中级以上职称海内外教学科研、博士后及企业研发人员；原文含在职博士生、全日制博士生推荐途径，适用细节须按申请人类别确认；校内立项原则不超过50%。','项目组原则上含1名实验室全职科研人员；具体到访天数未载。',
 '科研直接支出；原则不向外单位直接外拨，合法票据寄室报销；按学校纵向经费办法；禁止基建、偿债、罚款、捐赠、投资等无关支出。','一般1年；重点1—2年。',
 '一般SCI/EI-JA/CSCD核心非扩展论文1篇；重点至少1件申请并授权发明专利及软硬件可演示系统、源代码，专利可结题后1年补授权。总结3份、经费结算1份、成果复印1份；校外作者允许所在单位第一署名，同时含实验室署名及基金编号。',
 '负责人享有著作权及发表权（保密论文除外）；成果实验室与依托单位共享。','2022-10-04—11-04；签章原件3份寄达及Word发fhhuang@ygu.edu.cn；福州马尾登龙路99号30-304，350015。'
 ],'官网正文已完整取得；2026—2027指南未检出；2022申请附件未取得，历史机制不表示当期受理。','黄风华0591-83976047；13705073724；fhhuang@ygu.edu.cn')
add('N007','西南石油大学','机器人工程与智能制造南充市重点实验室—油气装备教育部重点实验室联合开放基金','P2 四川·南充','2026','','官方评审公示含蛇形机器人重大装备深腔检测等；原始任务指南未取得','公示12项，不能由立项数推算单项金额','https://www.swpu.edu.cn/jdy/info/1040/32450.htm',2,[unknown]*7,'仅取得2026-04-02拟立项公示，非申报通知。需原2026联合指南：任务、截止、资格合作、金额、周期、外拨、成果和知识产权。')
add('N008','科大讯飞股份有限公司／中国科学技术大学','认知智能全国重点实验室2026年高等教育科研专项','P2 安徽·合肥','2026','2026-08-15','高教教学管评测智能体；具身智能场景感知及虚实融合数据处理，机器人具体任务适用范围待核','已取得正文未载','https://cogskl.iflytek.com/archives/3471',2,[unknown,unknown,unknown,'不超过2年。',unknown,unknown,'电子申报至jfjiang6@iflytek.com，截止2026-08-15；初审后按通知交双面签章纸质2份。'],'须核高教专项具身方向具体任务与机器人联系，以及管理办法资格、金额、外拨、署名和权属；属于已截止线索，不是当前受理。','江老师；jfjiang6@iflytek.com')
add('N009','中国科学院宁波材料技术与工程研究所','浙江省机器人与智能制造装备技术重点实验室长期开放机制','P2 浙江·宁波','历史机制／年度待核','','精密驱动与控制、智能感知与测试、工业机器人；官方介绍明确设开放基金和开放课题','年度金额未取得','https://robotics-lab.nimte.ac.cn/view-13552.html',2,[unknown]*7,'机制有官方证据，但2026—2027正式年度指南未取得；资格、合作、金额、窗口、经费、周期、成果及权属均需当期文件。','郑天江0574-86324587')
add('N010','新疆大学','新疆农牧机器人及智能装备工程研究中心长期开放机制','P2 新疆·乌鲁木齐','历史机制／年度待核','','农牧机器人与智能装备；取得开放课题管理制度，年度任务未取得','中心建设总经费不是单项额度；年度金额未取得','https://jxxy.xju.edu.cn/__local/0/23/6F/B2938303B9580E6741F5D6F72BE_A1C96726_4AE3B.pdf',2,[unknown,'制度列客座交流及执行期至少1次学术报告；年度合作细则未取得。',unknown,unknown,unknown,unknown,unknown],'管理制度文本存在其他领域模板措辞，不能推填年度申报条件；须2026—2027指南及实际资格、单项金额、窗口、周期、成果权属和外拨确认。')
add('N011','辽宁石油化工大学','辽宁省石油化工智能优化与安全运行重点实验室','P2 辽宁·抚顺','2025','','年度指南线索为过程优化、安全运行、深度学习；机器人巡检见个人立项，不替代指南任务','年度额度未取得；个人立项3万元不直接套用','https://sice.lnpu.edu.cn/__local/6/85/AD/618AB99B1482BA4F1A20906037B_DEDEC0FE_32C8D.pdf',2,[
 '官方搜索片段：中级以上或博士，单位同意；同年申报或参与限1项。','片段：至少1名实验室固定人员参与。',unknown,'片段：2025-10-01—2026-09-30或2027-09-30，1—2年。',unknown,unknown,unknown],
 '官方PDF下载TLS失败；机器人具体年度任务、截止、金额及全部条款须读取原PDF。个人研究项目仅证明机制线索，不算新的年度公告。')
add('N012','吉林大学','工程仿生教育部重点实验室','P2 吉林·长春','2026','2026-04-28','松软地面机械仿生行走理论与技术；生物生产及加工机械设计，不将仿生材料自动算机器人任务','5万元/项；拟8项','https://bionic.jlu.edu.cn/info/1041/1693.htm',1,[
 '非吉林大学国内外高校/研究院人员；博士或副高，否则2名副高专家推荐；主持和参与共计≤2项；本室在研未结题负责人不得申报；需有团队成员。','必须有本室固定人员参加并任内部联系人；访问天数未载。',unknown,
 '正文称一般不超过2年，却要求填写2026-06-01—2027-05-30，约1年；以通知下达为准，须确认差异。',
 '二选一：申请发明1件加SCI/EI论文2篇；或核心论文3篇其中SCI/EI至少2篇。实验室为完成单位之一，联合署名并标基金与编号；结束交成果报告。','实验室与负责人单位共有；专利须含实验室申请单位。',
 '电子Word和PDF及负责人联系方式至xinyi@jlu.edu.cn；收到回复才算成功。2026-04-28 16:00截止。'
 ],'周期原文差异需确认；管理PDF验证码，另一官方渠道检索未取得全文。需管理办法开支、外拨、访问及申报模板。','辛老师0431-85095575／85095760-313；xinyi@jlu.edu.cn',time='16:00（北京时间）',sources=['https://bionic.jlu.edu.cn/system/_content/download.jsp?owner=1503487058&urltype=news.DownloadAttachUrl&wbfileid=18114730'])
add('N013','吉林大学','汽车底盘集成与仿生全国重点实验室','P2 吉林·长春','2026','2026-04-15','底盘动力学一体化自学习控制与域控制器；数字孪生智能运维、动力学仿真与测试装备','一般5—8万元；重点20万元','https://acib.jlu.edu.cn/info/1041/1643.htm',1,[
 '国内外汽车相关高校教师/企业科研人员；一般博士或中级，重点原则高级职称且博士。限项正文未载。','必须有实验室固定研究人员为合作者；保证时间，天数未载。',unknown,'一般3年。',
 '国际权威期刊论文、新标准、新技术原理样机等；统一数量、署名细则正文未载，不能套2025。',unknown,
 '截止2026-04-15；纸质2份顺丰寄吉林大学南岭校区人民大街5988号，130025；Word和PDF各1份至cuishan@jlu.edu.cn。'
 ],'管理文件尚未取得：经费开支/外拨、限项、成果数量/署名和IP；2026年度与M117的2025年度分开。','崔老师0431-85095090；13596182066；cuishan@jlu.edu.cn')
add('N014','江西飞行学院','低空地理信息与航路江西省教育厅重点实验室','P2 江西·南昌','2026','2026-08-30','低空航路规划；无人机AED快速投送与空地协同调度；空地协同城市航路规划','重点5万元；一般2万元；共6—8项','https://jcxy.jxfu.edu.cn/info/1042/14661.htm',1,[
 '一般≤55岁，高级职称或博士教学科研人员；每人限1项，不重复已资助/已有重大成果选题。','团队原则≤5人；依托本室一个课题组，获批后合作研究和经费报销。','与依托课题组合作报销；外拨、开支科目和管理费正文未载，管理条例未取得。','2年。',unknown,unknown,
 '申请表纸质3份寄送，电子至jxchang@jxfu.edu.cn；2026-08-30截止；南昌卧龙路269号330088。'
 ],'申报附件RAR下载超时；另官方渠道未取得条例。需指南任务细则、管理办法成果数量、IP、署名及经费外拨。','常老师15377677210；jxchang@jxfu.edu.cn',sources=['https://jcxy.jxfu.edu.cn/system/_content/download.jsp?owner=2063664427&urltype=news.DownloadAttachUrl&wbfileid=F1546CE1EE4272F24E211B6ABEFF1453'])
add('N015','湖南工学院','科技创新平台开放课题：汽车零部件产业链产教融合专项','P2 湖南·衡阳','2026','2026-07-10','方向6：轻型汽车电控悬架系统动力学建模与智能控制策略；其他育人管理方向不纳入','5万元/项；专项拟2项','https://www.hnit.edu.cn/kjc/info/1134/7693.htm',1,[
 '校内外博士或高级职称，否则两名高级推荐；相关研究基础，近3年相关项目或产业经验优先；主持限1项，参与不限；已主持2项省部级以上在研、研究精力不足者及本校开放课题未结題负责人不得申报（保留原文）。',
 '校外人员必须联合本校1名在职教师任合作责任人。',unknown,'1年。',unknown,unknown,
 '电子至1005105362@qq.com，主题2026产教融合专项－姓名－课题；A4双面纸质3份，普通封面直接装订至机械学院4215；2026-07-10 16:00截止。'
 ],'2026修订管理办法PDF验证码；另一官方渠道未取得。需经费开支/外拨、成果数量署名、IP及签章模板要求。','杜雪林15009295709；1005105362@qq.com',time='16:00（北京时间）',sources=['https://www.hnit.edu.cn/system/_content/download.jsp?owner=1438105372&urltype=news.DownloadAttachUrl&wbfileid=11958961'])

update('R006',1,status='已截止',kind='正式开放课题公告',date='2026-03-31',time='18:00（北京时间；纸质寄出4月3日18:00）',amount='30—100万元/项',amountlong='30—100万元/项',fields=[
 '国内外高校、科研院所、企业正式在职科研人员，原则博士或中级以上；单位书面同意；同一执行期仅申报/承担1项本基金。','正文未载必须固定合作者或具体到访天数，管理办法未取得。',
 '合同签订30%、中期评审50%、结题20%；开支、管理费及能否外拨未载，不能将拨付节点等同于外拨承诺。','1年：2026-05-01—2027-04-30。',
 '创新算法、文档、代码、软著/专利；CCF A或中科院一区及以上论文，数量未载。论文/报告要求实验室为第一作者/通讯作者/共同第一作者的第一单位；软件代码显著署名及基金号，不规范不计成果。',
 '论文、模型、代码、系统、数据、报告等共享；专利、软件需实验室共同申请。','Word与签章PDF 2026-03-31 18:00前至lihaipeng@gml.ac.cn；签章纸质1份须2026-04-03 18:00前寄出（以寄出为准）；深圳光明玉塘科润大厦B座10楼。'],
 gaps='首轮完整正文与申请书已取得，不套第二批M097；管理办法未取得：开支/外拨、管理费、访问天数待补；论文数量未载。网页发布时间与落款不同，保留证据，不改变截止。',contact='李海鹏18924592983；lihaipeng@gml.ac.cn',sources=['https://www.gml.ac.cn/xxgs/923.html','https://www.gml.ac.cn/core/extend/kindeditor/attached/file/20260226/20260226174223_19105.docx'])
update('M068',1,status='已截止',date='2026-10-08',time='具体时点未公开（日期已过）',amount='一般2万元；重点4万元（以评审为准）',fields=[
 '高级职称或相关学科博士；没有中级单独可申请条款。','成果须至少含1名本实验室固定研究人员，固定合作未落实；访问天数正文未载。','报销经办公室主任审核、主任和负责人共同签字；外拨和科目未载。','2年：2026-09-01—2028-08-31。',
 '重点≥1篇负责人一作/通讯的中科院大类SCI一区或CCF A；一般≥1篇SCI二区或CCF B；实验室必须第一完成单位，至少1固定人员作者；基金与编号标注必需；结题答辩。原文“第二年年末中期及下一年计划”与2年周期待确认。',unknown,
 '电子至hhy980344@163.com；纸质申请书2份，不能沿用旧3份；获批任务书2份共同签字；太原学院路3号软件实验大楼508。'],gaps='截至北京时间10月9日，10月8日截止已过；管理办法/模板未取得，外拨、IP、访问及中期时间条款待补。',contact='韩慧妍13703516841；hhy980344@163.com')
update('M043',1,status='已截止',date='2026-10-08',time='具体时点未公开（日期已过）',fields=[
 '国内外中级及以上职称科研人员；博士不替代职称；申请人同年限1项（不含参与者）。','固定成员联合申请优先，并非强制；至少1次到室学术交流、1次相关国内外学术会议。',unknown,'2026年10月—2028年10月。',
 '中北A2及以上论文至少1篇，或B及以上至少2篇；实验室第一标注单位，基金与编号致谢。',unknown,
 '2026-09-24—10-08；签章纸质3份至太原学院路3号，签章扫描电子至lnzhang0810@163.com。'],gaps='截止日期已过；管理办法及论文等级办法未取得，外拨、IP、具体A2/B认定需补。',contact='亓雪18845876210；lnzhang0810@163.com')
update('M035',1,status='已截止',date='2026-10-08',time='原文“10月8日前”，具体时点未载（日期已过）',gaps=pm['M035']['gaps']+'；北京时间已到10月9日，原截止日期已过。取得官方索引全文片段，直接访问404/TLS失败；具体金额和完整附件仍须手动补，不猜测。')
update('M041',fields=[
 '主要校外；博士或副高，年龄≤40；每人1项，单位≤2；课题组成员“不能同时参加两个以上项目”原文保留。','鼓励联合实验室成员；获批后确定/推荐校内联系人辅助报销；访问天数未载。','明确不外拨，实验室统一管理，依中科大科研经费办法报销；结题余款原则不再用于原课题。','1—2年申请人确定；接到资助通知起，1个月内签任务书。','每年进展及结题报告、成果证明；实验室署名和项目号，不标不计。正文未统一规定论文数量。','实验室与负责人单位共同所有。','单位签章纸质4份及电子至niant209@ustc.edu.cn；主题姓名-教育部重点实验室开放课题申请-课题名；2026-10-31截止。'],gaps='正文已取得完整资格/外拨/权属，不能重复列为未知；年度任务为脑成像、脑启发视觉听觉嗅觉、知识图谱推理决策、芯片系统。机器人相关具体任务适用范围仍待核；申请/任务书等附件未取得。',contact='童老师0551-63607515；niant209@ustc.edu.cn')

for id in ['M106','M114']:
 p=pm[id];remaining=p['gaps'].split('\n',1)[1] if '\n' in p['gaps'] else p['gaps']
 update(id,1,status='已截止',time='具体时点未公开（日期已过）',topic=p['topic'].replace('（截止日当天时点待确认）',''),gaps='北京时间10月9日，10月8日申报日期已过；既有已核研究任务和条款完整复用，不因本轮页面受限倒退证据。\n'+remaining)
for id in ['M068','M043','M035']:
 p=pm[id];p['topic']=p['topic'].replace('（截止日当天时点待确认）','');rebuild(p)
add('N016','天津大学','水利工程智能建设与运维全国重点实验室','P2 天津','2026','2026-01-31','少人化/无人化智能建设装备；异质智能建设装备群协同与规模化施工','资料未取得｜详细年度PDF验证码','https://www.tju.edu.cn/info/4732/81752.htm',1,[unknown,unknown,unknown,unknown,unknown,unknown,'2025-11-18—2026-01-31受理；hsskl@tju.edu.cn；具体材料要求见受限年度PDF，不能套天津大学其他实验室。'],'正式公告任务及受理截止已核，中文指南PDF验证码；另官方渠道未取得全文。需资格、固定合作、金额、周期、经费外拨/报销、成果IP及具体材料。','hsskl@tju.edu.cn',sources=['https://www.tju.edu.cn/system/_content/download.jsp?owner=2111843594&urltype=news.DownloadAttachUrl&wbfileid=0B6513F198BAEE0063A23A18FD934464'])

# Only actual newly checked exclusions are appended, not every irrelevant search hit.
for id,org,lab,reason,url in [
 ('E062','哈尔滨工业大学','2026海上风电叶片检修机器人专项','正式公告明确负责人为我校全职教师，校外负责人不符合；2026-09-24截止。','https://robot.hit.edu.cn/2026/0917/c338a401687/page.htm'),
 ('E063','哈尔滨工业大学','2026高压环境机器人专项','正式公告明确负责人为我校全职教师，校外负责人不符合；2026-09-24截止。','https://robot.hit.edu.cn/2026/0917/c338a401682/page.htm'),
 ('E064','东南大学','移动通信全国重点实验室2027','全年指南任务为无线网络、通信传输/信道、射频与基带芯片、光通信；无具体机器人相关任务，不能因智能组网保留。','https://ncrl.seu.edu.cn/2026/0803/c15598a578768/page.htm'),
 ('E065','昆明理工大学','复杂系统与类脑智能省重点实验室2026','具体任务是脑功能网络分析、认知行为及高阶传播模型/群体状态相变；未据非线性协同推为机器人控制。','https://lxy.kust.edu.cn/info/1151/2522.htm')]:
 D['excluded'].append([id,org,lab,'剔除（本轮新增）',reason,url,'2026-10-09'])

rank=['P0 南京及周边','P0 澳门','P0 香港','P0 广州及周边','P1 北京及周边','P1 西安及周边','P1 重庆']
for g in D['groups']:g['ids'].sort(key=lambda id:(rank.index(pm[id]['region']) if pm[id]['region'] in rank else 7,pm[id]['date'] or '9999-99-99',pm[id]['org'],id))
D['stats']={k:len(g['ids']) for k,g in zip(['current','historical','pending','other'],D['groups'])};D['stats'].update(total=len(D['projects']),excluded=len(D['excluded']))
D['national_changes']={'new_ids':new,'updated_ids':changes,'input_total':150,'execution_dates':'2026-10-08—09','deadline_baseline':'2026-10-09（北京时间）'}
(H/'prepared_regions.json').write_text(json.dumps(D,ensure_ascii=False,indent=2),encoding='utf8')
print(json.dumps({'stats':D['stats'],'changes':D['national_changes']},ensure_ascii=False))
