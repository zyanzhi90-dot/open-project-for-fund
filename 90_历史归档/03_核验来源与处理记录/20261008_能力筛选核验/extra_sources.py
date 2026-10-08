import pathlib,json,concurrent.futures,urllib.request,ssl,re
from collect import htmlparse
H=pathlib.Path(__file__).parent
extra={
 'M166':['https://ihealth.ccnu.edu.cn/xwzx/content.jsp?urltype=news.NewsContentUrl&wbnewsid=1811&wbtreeid=1491','https://ihealth.ccnu.edu.cn/system/_content/download.jsp?owner=2138426174&urltype=news.DownloadAttachUrl&wbfileid=B348CF146C1E3F71608EE3472FD5EE73'],
}
def retry(f):
 d=json.loads(f.read_text(encoding='utf8'))
 if d['id']=='M193':return
 urls=extra.get(d['id'],[])
 if not any(s['live'] for s in d['sources']) and any('CERTIFICATE_VERIFY' in a['status'] for a in d['attempts']):urls.append(d['url'])
 for url in urls:
  try:
   with urllib.request.urlopen(urllib.request.Request(url,headers={'User-Agent':'Mozilla/5.0'}),context=ssl._create_unverified_context(),timeout=15) as r:b=r.read();ct=r.headers.get('Content-Type','')
   if b[:2]==b'PK':
    from docx import Document
    import io
    doc=Document(io.BytesIO(b));text='\n'.join(p.text for p in doc.paragraphs)+'\n'+'\n'.join(' | '.join(c.text for c in row.cells) for t in doc.tables for row in t.rows);kind='官方FAQ附件'
   else:
    text,links=htmlparse(b.decode('utf8',errors='replace'),url);kind='本轮官方正文补核'
   if len(text)>100 and not re.search('无权访问|验证码|404|Access Denied',text[:1000]):d['sources'].insert(0,{'url':url,'kind':kind,'text':text,'live':True});d['attempts'].append({'url':url,'status':'补核全文已读取'})
   else:d['attempts'].append({'url':url,'status':'补核仍访问受限'})
  except Exception as e:d['attempts'].append({'url':url,'status':str(e)[:180]})
 if urls:f.write_text(json.dumps(d,ensure_ascii=False,indent=2),encoding='utf8');print(d['id'],d['attempts'][-1]['status'],flush=True)
with concurrent.futures.ThreadPoolExecutor(max_workers=8) as pool:list(pool.map(retry,(H/'sources').glob('*.json')))
