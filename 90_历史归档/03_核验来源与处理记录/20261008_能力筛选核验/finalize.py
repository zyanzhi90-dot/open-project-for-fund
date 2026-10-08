import pathlib,json,zipfile,xml.etree.ElementTree as E,hashlib,datetime,re
from openpyxl import load_workbook
H=pathlib.Path(__file__).parent;ROOT=H.parents[2]
D=json.loads((H/'prepared.json').read_text(encoding='utf8'));I=json.loads((H/'input.json').read_text(encoding='utf8'))
NS='http://schemas.openxmlformats.org/spreadsheetml/2006/main';REL='http://schemas.openxmlformats.org/officeDocument/2006/relationships';PKG='http://schemas.openxmlformats.org/package/2006/relationships'
E.register_namespace('',NS);E.register_namespace('r',REL)
def tag(t):return '{'+NS+'}'+t
def serialize(x):return E.tostring(x,encoding='utf-8',xml_declaration=True)
with zipfile.ZipFile(H/'authored.xlsx') as z:files={n:z.read(n) for n in z.namelist()}
def addlinks(num,links):
 name=f'xl/worksheets/sheet{num}.xml';rname=f'xl/worksheets/_rels/sheet{num}.xml.rels'
 sh=E.fromstring(files[name]);old=sh.find(tag('hyperlinks'))
 if old is not None:sh.remove(old)
 node=E.Element(tag('hyperlinks'))
 rel=E.fromstring(files[rname]) if rname in files else E.Element('{'+PKG+'}Relationships')
 ids={r.attrib['Id'] for r in rel};i=1000
 for cell,url,location in links:
  attr={'ref':cell}
  if location:attr['location']=location
  else:
   while 'rId'+str(i) in ids:i+=1
   rid='rId'+str(i);ids.add(rid);i+=1
   E.SubElement(rel,'{'+PKG+'}Relationship',{'Id':rid,'Type':REL+'/hyperlink','Target':url,'TargetMode':'External'})
   attr['{'+REL+'}id']=rid
  E.SubElement(node,tag('hyperlink'),attr)
 # OOXML schema requires hyperlinks before print settings and tableParts.
 later={'printOptions','pageMargins','pageSetup','headerFooter','rowBreaks','colBreaks','customProperties','cellWatches','ignoredErrors','smartTags','drawing','legacyDrawing','legacyDrawingHF','picture','oleObjects','controls','webPublishItems','tableParts','extLst'}
 idx=next((i for i,x in enumerate(sh) if x.tag.split('}')[-1] in later),len(sh));sh.insert(idx,node)
 files[name]=serialize(sh);files[rname]=serialize(rel)
links=[]
for idx,p in enumerate(D['projects'],2):
 links.extend([(f'C{idx}',None,f"'02项目完整详情'!A{idx+5}"),(f'F{idx}',p['url'],None)])
addlinks(1,links)
addlinks(2,[(f'Q{i}',p['url'],None) for i,p in enumerate(D['projects'],7)])
for num in [4]:
 addlinks(num,[(f'D{i}',r[3].splitlines()[0],None) for i,r in enumerate(D['logs'],2) if r[3].startswith('http')])
detail=E.fromstring(files['xl/worksheets/sheet2.xml']);cols=detail.find(tag('cols'))
if cols is not None:
 for c in cols:
  a=int(c.attrib['min']);b=int(c.attrib['max'])
  if 23<=a<=b<=32 or a==b==37:c.set('hidden','1')
files['xl/worksheets/sheet2.xml']=serialize(detail)
book=E.fromstring(files['xl/workbook.xml']);book.find(tag('sheets'))[-1].set('state','hidden');files['xl/workbook.xml']=serialize(book)
out=ROOT/'全国机器人开放课题_个人申报台账_筛选核验版_20261008.xlsx'
with zipfile.ZipFile(out,'w',compression=zipfile.ZIP_DEFLATED) as z:
 for n,b in files.items():z.writestr(n,b)
# Read-only validation against the input and authoritative prepared data.
src=ROOT/'全国机器人开放课题_个人申报台账_六列展示版_20261008.xlsx'
assert hashlib.sha256(src.read_bytes()).hexdigest().upper()==D['source_hash'].upper()
wb=load_workbook(out,data_only=False);home=wb['01申报总览'];detail=wb['02项目完整详情'];archive=wb['90筛选前原始记录']
assert home.max_column==6 and home.max_row==153
assert len({p['id'] for p in D['projects']})==152
assert {p['id'] for p in D['projects']}|{x[0] for x in D['excluded']}=={r[0] for r in I['rows']}
assert not ({p['id'] for p in D['projects']}&{x[0] for x in D['excluded']})
assert home.freeze_panes=='A2' and detail.freeze_panes=='E7'
assert next(iter(home.tables.values())).ref=='A1:F153'
assert archive.sheet_state=='hidden' and archive.max_row==200
for i,p in enumerate(D['projects'],2):
 r=i+5
 assert home.cell(i,1).value==p['status'] and home.cell(i,5).value==p['amount']
 assert detail.cell(r,1).value==p['id']
 assert home.cell(i,3).hyperlink.location==f"'02项目完整详情'!A{r}"
 assert home.cell(i,6).hyperlink.target==p['url']==detail.cell(r,17).hyperlink.target
 if p['date']:
  assert home.cell(i,2).value.strftime('%Y-%m-%d')==p['date']==detail.cell(r,5).value.strftime('%Y-%m-%d')
 else:assert home.cell(i,2).value=='待核'
 for col,v in enumerate(p['details'],1):
  if isinstance(v,str):v=re.sub(r'[\x00-\x08\x0b\x0c\x0e-\x1f]','',v)
  if col!=5:assert detail.cell(r,col).value==v or (v in [None,''] and detail.cell(r,col).value is None),(p['id'],col,detail.cell(r,col).value,v)
assert [list(r) for r in wb['03机构检索覆盖'].values]==D['coverage']
assert len(D['coverage'])==len(I['raw']['03机构检索覆盖'])
assert {r[1] for r in D['coverage'][6:]}=={r[1] for r in I['raw']['03机构检索覆盖'][6:]}
for s in wb:
 for row in s:
  for c in row:assert c.data_type!='e',(s.title,c.coordinate,c.value)
for i,r in enumerate(I['rows'],2):
 assert [archive.cell(i,j).value for j in range(1,33)]==r
seen=set();last=None
for p in D['projects']:
 if p['org_key']!=last:
  assert p['org_key'] not in seen
  seen.add(p['org_key']);last=p['org_key']
report={'output':str(out),'stats':D['stats'],'input_unchanged':True,'home_and_details_consistent':True,'original_199_preserved':True,'institution_grouping_valid':True,'native_filter_freeze_links_valid':True,'sha256':hashlib.sha256(out.read_bytes()).hexdigest()}
(H/'validation.json').write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf8')
print(json.dumps(report,ensure_ascii=False))
