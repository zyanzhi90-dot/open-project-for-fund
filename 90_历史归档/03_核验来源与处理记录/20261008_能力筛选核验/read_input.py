import pathlib,json,hashlib,sys
import openpyxl
from pypdf import PdfReader
ROOT=pathlib.Path(__file__).resolve().parents[3]
HERE=pathlib.Path(__file__).parent
OLD=HERE.parent/'rebuild'
src=ROOT/'全国机器人开放课题_个人申报台账_六列展示版_20261008.xlsx'
book=openpyxl.load_workbook(src,data_only=False)
raw={s.title:[list(r) for r in s.values] for s in book}
norm={x['id']:x for x in json.loads((OLD/'normalized.json').read_text(encoding='utf8'))['projects']}
rows=raw[book.sheetnames[1]][6:]
data={'source':str(src),'hash':hashlib.sha256(src.read_bytes()).hexdigest(),'raw':raw,'rows':rows,'norm':norm}
(HERE/'input.json').write_text(json.dumps(data,ensure_ascii=False,indent=2,default=lambda x:x.isoformat()),encoding='utf8')
pdf=ROOT/'面上项目-正文-2026-修改版-3.18-1.pdf'
pages=PdfReader(pdf).pages
text='\n'.join(f'=== PAGE {i+1} ===\n'+(p.extract_text() or '') for i,p in enumerate(pages))
(HERE/'proposal_full.txt').write_text(text,encoding='utf8')
if len(sys.argv)>1:
 ids=sys.argv[1].split(',')
 for r in rows:
  if r[0] in ids:
   f=OLD/'official'/(r[0]+'.json');p=json.loads(f.read_text(encoding='utf8')) if f.exists() else {}
   print('\n###',r[0],r[2],r[3],'\n',p.get('text','')[:13500])
else:print(json.dumps({'rows':len(rows),'pdf_pages':len(pages),'sheets':book.sheetnames},ensure_ascii=False))
