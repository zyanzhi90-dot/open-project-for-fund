import json,pathlib,collections,re,zipfile
from datetime import datetime
import openpyxl
root=pathlib.Path(__file__).parent
d=json.load(open(root/'normalized.json',encoding='utf8'))
rowmap=json.load(open(root/'detail_row_map.json',encoding='utf8'))
path=next((root.parent.parent/'outputs'/'20261008_robot_fund_rebuild').glob('*.xlsx'))
w=openpyxl.load_workbook(path,data_only=False)
h,o,detail,c=w.worksheets[:4]
assert w.sheetnames[:4]==['01 重点申报','02 地域项目总览','03 项目完整详情','04 机构检索进度']
overview_rows=[r for r in range(5,o.max_row+1) if o.cell(r,2).value]
assert len(overview_rows)==199
ids=[o.cell(r,2).value for r in overview_rows]
original_ids=[r[0] for r in d['raw']['02项目完整详情'][6:] if r[0]]
assert set(ids)==set(original_ids) and len(set(ids))==199
count=collections.Counter(o.cell(r,3).value[0] for r in overview_rows)
assert dict(count)==d['counts']
ps={p['id']:p for p in d['projects']}
homeids=[]
for r in range(5,h.max_row+1):
 if not h.cell(r,4).value:continue
 id=h.cell(r,4).value.split('\n')[-1];p=ps[id];homeids.append(id)
 assert p['grade'] in ['A','B'] and p['homepage'] and p['status_rank']<3
 assert h.cell(r,2).value==p['grade']+'类'
 assert h.cell(r,3).value==p['status']
 date=h.cell(r,5).value
 assert (date.strftime('%Y-%m-%d') if date else None)==p['date']
 assert h.cell(r,9).hyperlink.target==p['url']
 assert h.cell(r,4).hyperlink.location=="'03 项目完整详情'!A"+str(rowmap[id])
 assert detail.cell(rowmap[id],1).value==id
assert homeids==[p['id'] for p in d['home']]
seen=collections.defaultdict(dict)
for r in range(5,detail.max_row+1):seen[detail.cell(r,1).value][detail.cell(r,4).value]=detail.cell(r,5).value
for id,p in ps.items():
 assert seen[id]['科研匹配等级']==p['grade']
 assert seen[id]['申报状态']==p['status']
 for label,key in [('资助金额/类型','amount'),('申请人资格与限项','qualification'),('固定人员/实质合作要求','collaboration'),('经费外拨/使用范围','funds'),('执行周期','period'),('结题成果与署名','outputs'),('知识产权归属','ip'),('材料/签章/提交方式','materials')]:assert seen[id][label]==p[key]
assert sum(bool(c.cell(r,2).value) for r in range(5,c.max_row+1))==139
assert sum(x['original_count'] for x in d['coverage'])==144
for k,name in enumerate(['01申报总览','02项目完整详情','03机构检索覆盖']):
 s=w.worksheets[k+4]
 for ri,row in enumerate(d['raw'][name],1):
  for ci,v in enumerate(row,1):
   actual=s.cell(ri,ci).value
   if isinstance(v,str) and v.startswith('='):assert actual in [v,"'"+v]
   else:assert actual==v,(name,ri,ci,v,actual)
errors=[];formulas=0;hyperlinks=0
for s in w:
 for row in s.iter_rows():
  for cell in row:
   if cell.data_type=='e':errors.append((s.title,cell.coordinate,cell.value))
   if cell.data_type=='f':formulas+=1
   if cell.hyperlink:hyperlinks+=1
 assert len(s.tables)==1
 assert next(iter(s.tables.values())).autoFilter is not None
 assert s.freeze_panes is not None
assert not errors,errors
assert formulas==0
report={'records':len(ids),'grades':dict(count),'homepage':homeids,'coverage_original':144,'coverage_normalized':139,'native_hyperlinks':hyperlinks,'formula_errors':errors,'raw_cells_preserved':True,'sheets':[{'name':s.title,'rows':s.max_row,'cols':s.max_column,'freeze':s.freeze_panes,'filter':next(iter(s.tables.values())).ref} for s in w]}
(root/'qa_report.json').write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf8')
print(json.dumps(report,ensure_ascii=False,indent=2))
