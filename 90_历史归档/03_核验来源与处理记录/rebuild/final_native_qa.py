import pathlib,json,zipfile,collections,re
from lxml import etree as E
import openpyxl
R=pathlib.Path(__file__).parent;W=R.parents[1]
f=W/'outputs/20261008_robot_fund_audit/全国机器人开放课题_个人申报台账_信息核验版_20261008.xlsx'
links=json.load(open(R/'audit_all/final_links.json',encoding='utf8'))
NS='http://schemas.openxmlformats.org/spreadsheetml/2006/main';REL='http://schemas.openxmlformats.org/package/2006/relationships';RID='http://schemas.openxmlformats.org/officeDocument/2006/relationships'
with zipfile.ZipFile(f) as z:parts={n:z.read(n) for n in z.namelist()}
book=E.fromstring(parts['xl/workbook.xml']);rels=E.fromstring(parts['xl/_rels/workbook.xml.rels']);mp={x.get('Id'):x.get('Target') for x in rels};sheets={}
for s in book.findall('.//{'+NS+'}sheet'):
 t=mp[s.get('{'+RID+'}id')];sheets[s.get('name')]=('xl/'+t if not t.startswith('/') else t[1:])
for name,path in sheets.items():
 sl=[x for x in links if x['sheet']==name]
 if not sl:continue
 sh=E.fromstring(parts[path]);existing=sh.find('{'+NS+'}hyperlinks')
 if existing is not None:sh.remove(existing)
 hp=E.Element('{'+NS+'}hyperlinks');rn=str(pathlib.PurePosixPath(path).parent/'_rels'/(pathlib.PurePosixPath(path).name+'.rels'))
 rr=E.fromstring(parts[rn]) if rn in parts else E.Element('{'+REL+'}Relationships',nsmap={None:REL})
 for x in list(rr):
  if x.get('Type','').endswith('/hyperlink'):rr.remove(x)
 for i,l in enumerate(sl):
  id='auditLink'+str(i+1);E.SubElement(hp,'{'+NS+'}hyperlink',ref=l['cell'],attrib={'{'+RID+'}id':id})
  E.SubElement(rr,'{'+REL+'}Relationship',Id=id,Type=RID+'/hyperlink',Target=l['url'],TargetMode='External')
 # Preserve SpreadsheetML order: hyperlinks follow conditionalFormatting and precede print settings.
 late={'printOptions','pageMargins','pageSetup','headerFooter','rowBreaks','colBreaks','customProperties','cellWatches','ignoredErrors','smartTags','drawing','legacyDrawing','legacyDrawingHF','drawingHF','picture','oleObjects','controls','webPublishItems','tableParts','extLst'}
 index=next((i for i,x in enumerate(sh) if E.QName(x).localname in late),len(sh));sh.insert(index,hp)
 parts[path]=E.tostring(sh,xml_declaration=True,encoding='UTF-8',standalone=True);parts[rn]=E.tostring(rr,xml_declaration=True,encoding='UTF-8',standalone=True)
temp=f.with_suffix('.patched.xlsx')
with zipfile.ZipFile(temp,'w',zipfile.ZIP_DEFLATED) as z:
 for n,b in parts.items():z.writestr(n,b)
temp.replace(f)
raw=json.load(open(R/'original.json',encoding='utf8'));data=json.load(open(R/'audit_all/audited_original_layout.json',encoding='utf8'));w=openpyxl.load_workbook(f,data_only=False)
assert w.sheetnames==list(data)
orig={x[0]:x for x in raw['02项目完整详情'][6:]};detail={row[0].value:[c.value for c in row] for row in list(w['02项目完整详情'].rows)[6:]};summary={row[10].value:[c.value for c in row] for row in list(w['01申报总览'].rows)[6:]}
assert len(detail)==len(summary)==len(orig)==199
assert set(detail)==set(summary)==set(orig)
for id,row in detail.items():
 assert row[22:26]==orig[id][22:26],id+' historical fields changed'
 sr=summary[id]
 assert sr[2]==row[5] and sr[4]==row[4] and sr[6]==row[9] and sr[9]==row[16],id+' inconsistent summary'
 for c in [10,11,12,13,14,15,18,19,20,21]:assert row[c] is not None,id+' blank field'
assert [x[6].value for x in list(w['03机构检索覆盖'].rows)[6:] if x[1].value].count('本轮定向检索完成；不等于所有平台穷尽核查')==144
errors=[];nh=0
for sh in w:
 for row in sh:
  for c in row:
   if c.data_type=='e':errors.append((sh.title,c.coordinate,c.value))
   if c.hyperlink:nh+=1
 assert sh.freeze_panes,sh.title+' no freeze'
assert not errors,errors
assert nh>=len(links)
assert detail['M004'][4]=='2026-07-19'
assert detail['M047'][2]=='哈尔滨理工大学'
assert detail['M098'][4]=='2026-08-15' and '专利权人须为研究院' in detail['M098'][14]
assert detail['M107'][4]=='2026-10-12'
assert detail['M117'][4]=='2025-03-31'
assert detail['M165'][4]=='2025-08-15'
assert detail['M185'][9].startswith('本年度每项4—6万元')
assert detail['M194'][4]=='2026-05-06'
for id in ['M006','M064','M071','M166','M182']:assert detail[id][4]=='未确认正式截止日期'
rep={'projects_original_retained':199,'institutions_targeted_checked':144,'summary_consistent':True,'original_history_fields_retained':True,'native_hyperlinks':nh,'sheets':w.sheetnames,'errors':errors,'file':str(f),'bytes':f.stat().st_size}
(R/'audit_all/final_qa_report.json').write_text(json.dumps(rep,ensure_ascii=False,indent=2),encoding='utf8');print(json.dumps(rep,ensure_ascii=False))
