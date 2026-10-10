from pathlib import Path
from openpyxl import load_workbook
import json,re,datetime,hashlib
H=Path(__file__).parent;R=H.parents[2];source=R/'全国机器人开放课题_个人申报台账_结构精简版_20261009.xlsx'
w=load_workbook(source);detail=w['02项目完整详情'];today=datetime.date(2026,10,10)
rows=[list(r) for r in detail.values];pm={r[0]:{'row':i+1,'values':r} for i,r in enumerate(rows[1:],1)}
assert len(pm)==397
lists={};mapping=[]
for name in ['01申报总览','03历史项目','04待核项目','05其他资助']:
 s=w[name];lists[name]=[]
 for n in range(2,s.max_row+1):
  values=[c.value for c in s[n]];link=s.cell(n,3).hyperlink
  assert link and link.location
  rr=int(re.search(r'A(\d+)',link.location)[1]);p=rows[rr-1];ident=p[0]
  assert values[2]==p[2]+'·'+p[3] and values[1]==p[6],(ident,name,n)
  entry={'id':ident,'source_sheet':name,'source_row':n,'detail_row':rr,'url':s.cell(n,6).hyperlink.target if s.cell(n,6).hyperlink else p[19],'values':values[:5]+['公告']}
  lists[name].append(entry);mapping.append(ident)
assert len(mapping)==len(set(mapping))==397
def rank(p):
 v=p['values'];region=str(v[24] or '');org=str(v[2] or '')
 if 'P0 南京' in region or any(x in region+' '+org for x in ['南京','滁州']):return 0
 if '澳门' in region+' '+org:return 1
 if '香港' in region+' '+org:return 2
 if any(x in region+' '+org for x in ['广州','深圳','珠海','佛山','东莞','中山','惠州','珠三角']):return 3
 if '北京' in region+' '+org:return 4
 if '西安' in region+' '+org:return 5
 if '重庆' in region+' '+org:return 6
 return 7
def key(e):
 p=pm[e['id']];dt=p['values'][6]
 return rank(p),dt.isoformat() if isinstance(dt,datetime.datetime) else '9999-99-99',str(p['values'][2]),str(p['values'][3]),e['id']
home=lists['01申报总览'];hist=lists['03历史项目'];others=lists['05其他资助'];only=[];moves=[]
# Explicit existing historical/award facts; no inferred cutoff dates.
known_history={'N038','N039','N043','N050','N065','N063','N061','N068','N074','N075','N086','N087','N079','N101','N110','N138'}
for e in lists['04待核项目']:
 v=pm[e['id']]['values'];dt=v[6]
 expired=isinstance(dt,datetime.datetime) and dt.date()<today
 if expired or e['id'] in known_history:
  hist.append(e);moves.append({'id':e['id'],'from':'04待核项目','to':'03历史项目','reason':'明确截止日期已过' if expired else '既有记录已证实历史公告/立项/研究期结束；不补造具体截止日'})
 else:only.append(e)
patches=[]
for e in home+hist+others+only:
 v=pm[e['id']]['values'];old=v[4]
 if e in home:status='受理中'
 elif e in hist:
  status='已截止' if isinstance(v[6],datetime.datetime) and v[6].date()<today else '已结束·准确截止未载' if e['id'] in {'N038','N050','N079','N110'} else '历史记录·截止见详情'
 elif e in others:status=old
 elif isinstance(v[6],datetime.datetime):status='今日截止·时点未载' if v[6].date()==today else '受理窗口内·任务范围未明确'
 else:status='未发布正式申报公告' if e['id'] in {'M182','M064'} else '当期受理窗口未核实'
 if status!=old:patches.append({'row':pm[e['id']]['row'],'col':'E','before':old,'value':status});v[4]=status
 e['values'][0]=status
 # Remove links to the deleted sheet from display prose; preserve substantive clauses.
 for j in [18,23,26]:
  oldtext=v[j]
  if isinstance(oldtext,str) and any(x in oldtext for x in ['待核页','待核非首页']):
   newtext=oldtext.replace('仍在待核页','仅保留在完整详情').replace('待核页','完整详情').replace('待核非首页','仅保留详情，未列首页')
   patches.append({'row':pm[e['id']]['row'],'col':chr(65+j) if j<26 else 'AA','before':oldtext,'value':newtext});v[j]=newtext
home.sort(key=key);hist.sort(key=key)
assert all(isinstance(e['values'][1],datetime.datetime) and e['values'][1].date()>today for e in home)
def serial(v):return v.isoformat() if isinstance(v,(datetime.datetime,datetime.date)) else v
out={'source':str(source),'sha256':hashlib.sha256(source.read_bytes()).hexdigest(),'output_name':'全国机器人开放课题_个人申报台账_分类排序调整版_20261010.xlsx','home':home,'history':hist,'others':others,'detail_only':only,'moves':moves,'patches':patches,'details':list(pm.values()),'stats':{'全部详情':397,'首页正式受理':len(home),'历史项目':len(hist),'其他资助':len(others),'仅保留详情':len(only),'原待核转历史':len(moves)},'new_sheets':['01申报总览','02项目完整详情','03历史项目','04其他资助','05机构检索覆盖']}
(H/'plan.json').write_text(json.dumps(out,ensure_ascii=False,indent=2,default=serial),encoding='utf8')
print(json.dumps(out['stats'],ensure_ascii=False));print('ORDER',[(e['id'],rank(pm[e['id']]),str(e['values'][1])[:10]) for e in home]);w.close()
