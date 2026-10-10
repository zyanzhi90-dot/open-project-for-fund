from pathlib import Path
import json,zipfile,xml.etree.ElementTree as E,copy,posixpath,re,hashlib,datetime
from openpyxl import load_workbook
H=Path(__file__).parent;D=json.loads((H/'plan.json').read_text(encoding='utf8'))
S='http://schemas.openxmlformats.org/spreadsheetml/2006/main';Q='{'+S+'}';O='http://schemas.openxmlformats.org/officeDocument/2006/relationships';P='http://schemas.openxmlformats.org/package/2006/relationships'
E.register_namespace('',S);E.register_namespace('r',O)
def xml(x):
 # Package manifests require their own default namespace for Excel/renderer compatibility.
 E.register_namespace('',x.tag.split('}')[0].lstrip('{'))
 return E.tostring(x,encoding='utf-8',xml_declaration=True)
def unpack(path):
 with zipfile.ZipFile(path) as z:return {n:z.read(n) for n in z.namelist()}
files=unpack(D['source']);auth=unpack(H/'authored.xlsx')
def resolve(fn,target):return target.lstrip('/') if target.startswith('/') else posixpath.normpath(posixpath.join(posixpath.dirname(fn),target))
def sheetpaths(f):
 b=E.fromstring(f['xl/workbook.xml']);rels={x.get('Id'):resolve('xl/workbook.xml',x.get('Target')) for x in E.fromstring(f['xl/_rels/workbook.xml.rels'])}
 return {x.get('name'):rels[x.get('{'+O+'}id')] for x in b.find(Q+'sheets')}
origpaths=sheetpaths(files);authpaths=sheetpaths(auth)
strings=[]
if 'xl/sharedStrings.xml' in auth:strings=[''.join(x.itertext()) for x in E.fromstring(auth['xl/sharedStrings.xml'])]
def contents(c):
 c=copy.deepcopy(c);c.attrib.pop('s',None)
 if c.get('t')=='s':
  value=strings[int(c.find(Q+'v').text)];c.clear();c.set('t','inlineStr');is_=E.SubElement(c,Q+'is');E.SubElement(is_,Q+'t').text=value
 return c
def translated(ac,template,ref):
 c=contents(ac);c.set('r',ref)
 if template is not None and template.get('s') is not None:c.set('s',template.get('s'))
 return c
def rf(fn):return posixpath.join(posixpath.dirname(fn),'_rels',posixpath.basename(fn)+'.rels')
def replace_links(tree,fn,entries):
 old=tree.find(Q+'hyperlinks')
 if old is not None:tree.remove(old)
 relfile=rf(fn);rels=E.fromstring(files[relfile])
 for x in list(rels):
  if x.get('Type','').endswith('/hyperlink'):rels.remove(x)
 node=E.Element(Q+'hyperlinks')
 for n,e in enumerate(entries,2):
  E.SubElement(node,Q+'hyperlink',{'ref':f'C{n}','location':f"'02项目完整详情'!A{e['detail_row']}"})
  if e['url']:
   rid=f'rIdCallLink{n}';E.SubElement(node,Q+'hyperlink',{'ref':f'F{n}','{'+O+'}id':rid})
   E.SubElement(rels,'{'+P+'}Relationship',{'Id':rid,'Type':O+'/hyperlink','Target':e['url'],'TargetMode':'External'})
 later={'printOptions','pageMargins','pageSetup','headerFooter','rowBreaks','colBreaks','drawing','tableParts','extLst'}
 idx=next((i for i,c in enumerate(tree) if c.tag.split('}')[-1] in later),len(tree));tree.insert(idx,node)
 files[relfile]=xml(rels)
 return rels
for name,entries in [('04其他资助',D['others']),('03历史项目',D['history'])]:
 last='H' if name=='04其他资助' else 'F';fn=origpaths[name];tree=E.fromstring(files[fn]);sd=tree.find(Q+'sheetData');template=list(sd)[1];templates={c.get('r').rstrip('0123456789'):c for c in template}
 at=E.fromstring(auth[authpaths[name]]);arows={int(r.get('r')):r for r in at.find(Q+'sheetData')}
 for old in list(sd)[1:]:sd.remove(old)
 for n,e in enumerate(entries,2):
  row=E.Element(Q+'row',{**template.attrib,'r':str(n)})
  if arows[n].get('ht'):row.set('ht',arows[n].get('ht'));row.set('customHeight','1')
  for ac in arows[n]:
   ref=ac.get('r');col=ref.rstrip('0123456789');row.append(translated(ac,templates.get(col),ref))
  sd.append(row)
 dimension=tree.find(Q+'dimension')
 if dimension is not None:dimension.set('ref',f'A1:{last}{len(entries)+1}')
 rels=replace_links(tree,fn,entries)
 for rel in rels:
  if rel.get('Type','').endswith('/table'):
   tf=resolve(fn,rel.get('Target'));t=E.fromstring(files[tf]);t.set('ref',f'A1:{last}{len(entries)+1}');a=t.find(Q+'autoFilter')
   if a is not None:a.set('ref',t.get('ref'))
   files[tf]=xml(t)
 
 if name=='04其他资助':
  header=next(c for c in sd[0] if c.get('r')=='G1');ah=next(c for c in arows[1] if c.get('r')=='G1');idx=list(sd[0]).index(header);sd[0].remove(header);sd[0].insert(idx,translated(ah,header,'G1'))
 files[fn]=xml(tree)
fn=origpaths['02项目完整详情'];tree=E.fromstring(files[fn]);sr={int(r.get('r')):r for r in tree.find(Q+'sheetData')}
ad=E.fromstring(auth[authpaths['02项目完整详情']]);acells={c.get('r'):c for r in ad.find(Q+'sheetData') for c in r}
for p in D['patches']:
 ref=p['col']+str(p['row']);r=sr[p['row']];old=next((c for c in r if c.get('r')==ref),None);c=translated(acells[ref],old,ref)
 if old is not None:idx=list(r).index(old);r.remove(old);r.insert(idx,c)
 else:r.append(c)
files[fn]=xml(tree)

out=H/D['output_name']
with zipfile.ZipFile(out,'w',zipfile.ZIP_DEFLATED) as z:
 for fn,b in files.items():z.writestr(fn,b)
source=load_workbook(D['source']);w=load_workbook(out)
assert w.sheetnames==D['new_sheets']
patch={(p['row'],p['col']):p['value'] for p in D['patches']}
for r in source['02项目完整详情']:
 for c in r:
  nc=w['02项目完整详情'][c.coordinate]
  assert nc.value==patch.get((c.row,c.column_letter),c.value),(c.coordinate,'Out-of-scope value change')
  assert nc.style_id==c.style_id
  if c.hyperlink:assert nc.hyperlink==c.hyperlink
orig=unpack(D['source'])
for name in ['01申报总览','05机构检索覆盖']:
 assert files[origpaths[name]]==orig[origpaths[name]]
for name,entries in [('04其他资助',D['others']),('03历史项目',D['history'])]:
 s=w[name];last='H' if name=='04其他资助' else 'F'
 assert s.max_row==len(entries)+1 and s.freeze_panes=='A2' and len(s.tables)==1
 assert next(iter(s.tables.values())).ref==f'A1:{last}{s.max_row}'
 for n,e in enumerate(entries,2):
  for j,value in enumerate(e['values'],1):
   if j!=2:assert s.cell(n,j).value==value,(e['id'],j)
  assert s.cell(n,2).value==source[e['source_sheet']].cell(e['source_row'],2).value
  assert s.cell(n,3).hyperlink.location==f"'02项目完整详情'!A{e['detail_row']}"
  assert s.cell(n,6).hyperlink.target==e['url']
  dt=s.cell(n,2).value
  if isinstance(dt,datetime.datetime):
   assert (dt.date()<datetime.date(2026,10,10)) if name=='03历史项目' else (dt.date()>datetime.date(2026,10,10))
assert w['02项目完整详情'].max_row==398
assert {e['id'] for e in D['others']}=={'F002','F022'}
assert len({r[0] for r in list(w['02项目完整详情'].values)[1:]})==397
assert hashlib.sha256(Path(D['source']).read_bytes()).hexdigest()==D['sha256']
report={'output':str(out),'sha256':hashlib.sha256(out.read_bytes()).hexdigest(),'stats':D['stats'],'all_397_retained':True,'substantive_terms_and_sources_unchanged':True,'home_and_coverage_xml_unchanged':True,'links_filters_dates_verified':True,'visual_review':'pending','old_sha256':D['sha256']}
(H/'validation.json').write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf8')
print(json.dumps(report,ensure_ascii=False));source.close();w.close()
