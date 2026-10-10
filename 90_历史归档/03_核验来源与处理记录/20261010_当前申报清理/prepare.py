from pathlib import Path
from openpyxl import load_workbook
import json,re,datetime,hashlib
H=Path(__file__).parent; R=H.parents[2]
source=R/'全国机器人开放课题_个人申报台账_分类排序调整版_20261010.xlsx'
w=load_workbook(source); detail=w['02项目完整详情']; rows=list(detail.values)
pm={r[0]:{'row':i+2,'values':list(r)} for i,r in enumerate(rows[1:])}
def entries(name):
 out=[]; s=w[name]
 for n in range(2,s.max_row+1):
  rr=int(re.search(r'A(\d+)',s.cell(n,3).hyperlink.location)[1]); v=list(s[n]); ident=rows[rr-1][0]
  out.append({'id':ident,'source_sheet':name,'source_row':n,'detail_row':rr,'url':s.cell(n,6).hyperlink.target,'values':[c.value for c in v]})
 return out
home=entries('01申报总览'); hist=entries('03历史项目'); old=entries('04其他资助'); other=[]; only=[]; patches=[]; moves=[]
reasons={
 'F001':('任务范围未明确·仅留详情','正式受理日期已知，但已有公告未证实机器人相关具体任务；不以医疗器械名称代替任务证据。'),
 'F007':('尚未开放·2026-10-12开始','当前为2026-10-10，10月12日起受理；香港牵头及内地方资格见原条款。'),
 'F019':('广东单位牵头·省外参与资格未载','广东注册法人及广东在职负责人牵头；南京信息工程大学能否作为省外参与方、获得经费未载，不列当前个人申报入口。'),
 'F022':('滚动受理·省际合作参与','公告明确允许省外合作单位参与，须青海单位牵头；资金由协议确定，不把项目总额当作个人可获得金额。'),
 'F002':('受理中','公告有机器人/具身智能测控具体任务及高校团队申报入口；未公开条件继续保留。')}
for e in old:
 p=pm[e['id']]; dt=p['values'][6]; ident=e['id']
 if isinstance(dt,datetime.datetime) and dt.date()<datetime.date(2026,10,10):
  status='已截止'; reason='已有明确截止日期早于2026-10-10；停止补取历史资料。'; hist.append({**e,'values':e['values'][:5]+['公告']}); dest='03历史项目'
 elif ident in {'F002','F022'}:
  status,reason=reasons[ident]; other.append(e); dest='04其他资助'
 else:
  status,reason=reasons.get(ident,('当期受理窗口未核实·仅留详情','已有资料不能确定当期受理窗口或现金资助；不推定仍可申请。')); only.append(e); dest='02项目完整详情'
 oldstatus=p['values'][4]; patches.append({'row':p['row'],'col':'E','before':oldstatus,'value':status}); e['values'][0]=status
 if dest=='03历史项目':hist[-1]['values'][0]=status
 moves.append({'id':ident,'from':'04其他资助','to':dest,'reason':reason})
other[0]['values'][6]='高校团队申报；团队≥3人，具备独立硬件研究基础。现金一般5–25万元、重点25–100万元；平台支持另计。资格与结题细则见详情。'
other[1]['values'][4]='省外合作经费按协议确定（未公布）；联合项目总额上限300万元，不是个人额度'
other[1]['values'][6]='合作参与渠道：须青海单位牵头，省外第一合作单位参与；成熟技术成果（不含论文）及项目经历要求见详情。省外人数、外拨经费均≤50%；本人条件及合作落实待后续判断。'
for e in other:e['values'][5]='公告'
def rank(e):
 p=pm[e['id']]['values']; text=str(p[24] or '')+' '+str(p[2] or '')
 if 'P0 南京' in text or any(x in text for x in ['南京','滁州']):return 0
 for k,words in enumerate([['澳门'],['香港'],['广州','深圳','珠海','佛山','东莞','中山','惠州','珠三角'],['北京'],['西安'],['重庆']],1):
  if any(x in text for x in words):return k
 return 7
def key(e):
 dt=e['values'][1];return rank(e),dt.isoformat() if isinstance(dt,datetime.datetime) else '9999',e['values'][2],e['id']
hist.sort(key=key);other.sort(key=key)
oldlisted={e['id'] for e in home+entries('03历史项目')+old}
all_only=[x for ident,x in pm.items() if ident not in oldlisted]+only
stats={'全部详情':397,'首页正式受理':19,'历史项目':len(hist),'其他当前资助':len(other),'仅留详情':len(all_only),'其他资助转历史':11,'其他资助转详情':len(only)}
assert stats=={'全部详情':397,'首页正式受理':19,'历史项目':296,'其他当前资助':2,'仅留详情':80,'其他资助转历史':11,'其他资助转详情':9}
d={'source':str(source),'sha256':hashlib.sha256(source.read_bytes()).hexdigest(),'output_name':'全国机器人开放课题_个人申报台账_当前申报清理版_20261010.xlsx','history':hist,'others':other,'patches':patches,'moves':moves,'stats':stats,'new_sheets':w.sheetnames}
(H/'plan.json').write_text(json.dumps(d,ensure_ascii=False,indent=2,default=lambda v:v.isoformat()),encoding='utf8')
print(json.dumps(stats,ensure_ascii=False));w.close()
