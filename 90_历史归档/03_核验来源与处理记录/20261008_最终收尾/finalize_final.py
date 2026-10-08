import pathlib,json,zipfile,xml.etree.ElementTree as E,hashlib,datetime,re,collections
from openpyxl import load_workbook
H=pathlib.Path(__file__).parent;ROOT=H.parents[2];D=json.loads((H/'prepared_final.json').read_text(encoding='utf8'))
NS='http://schemas.openxmlformats.org/spreadsheetml/2006/main';REL='http://schemas.openxmlformats.org/officeDocument/2006/relationships';PKG='http://schemas.openxmlformats.org/package/2006/relationships'
E.register_namespace('',NS);E.register_namespace('r',REL)
def tag(t):return '{'+NS+'}'+t
def xml(x):return E.tostring(x,encoding='utf-8',xml_declaration=True)
with zipfile.ZipFile(H/'authored_final.xlsx') as z:files={n:z.read(n) for n in z.namelist()}
def links(num,seq):
 name=f'xl/worksheets/sheet{num}.xml';relname=f'xl/worksheets/_rels/sheet{num}.xml.rels';sh=E.fromstring(files[name]);old=sh.find(tag('hyperlinks'))
 if old is not None:sh.remove(old)
 node=E.Element(tag('hyperlinks'));rels=E.fromstring(files[relname]) if relname in files else E.Element('{'+PKG+'}Relationships');ids={x.attrib['Id'] for x in rels};i=1000
 for cell,url,loc in seq:
  attr={'ref':cell}
  if loc:attr['location']=loc
  elif url:
   while 'rId'+str(i) in ids:i+=1
   rid='rId'+str(i);ids.add(rid);i+=1
   E.SubElement(rels,'{'+PKG+'}Relationship',{'Id':rid,'Type':REL+'/hyperlink','Target':url,'TargetMode':'External'});attr['{'+REL+'}id']=rid
  else:continue
  E.SubElement(node,tag('hyperlink'),attr)
 later={'printOptions','pageMargins','pageSetup','headerFooter','rowBreaks','colBreaks','customProperties','cellWatches','ignoredErrors','smartTags','drawing','legacyDrawing','legacyDrawingHF','picture','oleObjects','controls','webPublishItems','tableParts','extLst'}
 idx=next((i for i,x in enumerate(sh) if x.tag.split('}')[-1] in later),len(sh));sh.insert(idx,node);files[name]=xml(sh);files[relname]=xml(rels)
detailmap={p['id']:i+2 for i,p in enumerate(D['projects'])};pm={p['id']:p for p in D['projects']}
for p in D['projects']:
 for a in p['aliases']:detailmap[a]=detailmap[p['id']]
seq=[]
for x in D['home_map']:
 p=pm[x['id']];r=x['row'];seq.extend([(f'C{r}',None,f"'02项目完整详情'!A{detailmap[p['id']]}"),(f'F{r}',p['url'],None)])
links(1,seq);links(2,[(f'T{i+2}',p['url'],None) for i,p in enumerate(D['projects'])]);links(4,[(f'D{i+2}',r[3].splitlines()[0],None) for i,r in enumerate(D['logs']) if r[3].startswith('http')]);links(5,[(f'F{i+2}',r[5],None) for i,r in enumerate(D['excluded'])])
book=E.fromstring(files['xl/workbook.xml'])
for s in book.find(tag('sheets')):
 if s.attrib['name'].startswith(('90','91')):s.set('state','hidden')
files['xl/workbook.xml']=xml(book)
out=ROOT/'全国机器人开放课题_个人申报台账_最终核验版_20261008.xlsx'
with zipfile.ZipFile(out,'w',compression=zipfile.ZIP_DEFLATED) as z:
 for n,b in files.items():z.writestr(n,b)
assert hashlib.sha256(pathlib.Path(D['source']).read_bytes()).hexdigest()==D['source_hash']
w=load_workbook(out,data_only=False);home=w['01申报总览'];detail=w['02项目完整详情'];arch=w['90筛选前原始记录']
assert home.max_column==6 and len(D['projects'])==166
assert home.freeze_panes=='A2' and detail.freeze_panes=='E2'
assert len(home.tables)==4 and detail.max_row==167 and arch.max_row==200
assert len(D['excluded'])==32 and len(detailmap)==167 and len(detailmap)+len(D['excluded'])==199
assert set(detailmap)&{x[0] for x in D['excluded']}==set()
assert {arch.cell(i,1).value for i in range(2,201)}==set(detailmap)|{x[0] for x in D['excluded']}
assert all(arch.cell(i,34).value==(pm[arch.cell(i,35).value]['status'] if arch.cell(i,35).value else next(x[3] for x in D['excluded'] if x[0]==arch.cell(i,1).value)) for i in range(2,201))
for x in D['home_map']:
 p=pm[x['id']];r=x['row'];dr=detailmap[p['id']]
 assert home.cell(r,1).value==detail.cell(dr,5).value==p['status']
 assert home.cell(r,3).hyperlink.location==f"'02项目完整详情'!A{dr}"
 assert home.cell(r,6).hyperlink.target==p['url']==detail.cell(dr,20).hyperlink.target
 if p['date']:assert home.cell(r,2).value.strftime('%Y-%m-%d')==detail.cell(dr,7).value.strftime('%Y-%m-%d')==p['date']
 assert home.cell(r,4).value==p['topic'] and home.cell(r,5).value==p['amount']
 rank={'可申请':0,'受理中·条件待核':1,'截止待确认':2,'线索待核':3,'已截止':4,'资格不符':5}
main=[p for p in D['projects'] if not p['id'].startswith('F')]
assert [rank[p['status']] for p in main]==sorted(rank[p['status']] for p in main)
for status in rank:
 seq=[p['org'] for p in main if p['status']==status];blocks=[]
 for org in seq:
  if not blocks or blocks[-1]!=org:blocks.append(org)
 assert len(blocks)==len(set(blocks)),(status,blocks)
for s in w:
 for row in s:
  for c in row:assert c.data_type!='e',(s.title,c.coordinate)
for s in [home,detail,w['03机构检索覆盖']]:
 for row in s:
  for c in row:
   if isinstance(c.value,str):assert not re.search('适合承担|可承担|申请书未显示能力|定向检索完成|无需重复',c.value),(s.title,c.coordinate)
assert [p['id'] for p in D['projects'][-4:]]==[p['id'] for p in D['projects'] if p['id'].startswith('F')]
report={'output':str(out),'stats':D['stats'],'input_unchanged':True,'original_199_covered':True,'stable_ids_links_dates_filters_freeze_valid':True,'status_order_valid':True,'sha256':hashlib.sha256(out.read_bytes()).hexdigest()}
(H/'validation_final.json').write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf8')
print(json.dumps(report,ensure_ascii=False))
