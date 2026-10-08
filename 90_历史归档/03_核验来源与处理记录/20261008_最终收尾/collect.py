import pathlib,json,re,urllib.request,urllib.parse,ssl,io,concurrent.futures,time,sys
from pypdf import PdfReader
from docx import Document
from html import unescape
H=pathlib.Path(__file__).parent;P=H.parent/'20261008_能力筛选核验';OLD=H.parent/'rebuild'
D=json.loads((P/'input.json').read_text(encoding='utf8'));N=D['norm'];(H/'sources').mkdir(exist_ok=True)
def htmlparse(body,url):
 body=re.sub(r'<(script|style)\b[^>]*>.*?</\1>','',body,flags=re.I|re.S)
 links=[]
 for href,label in re.findall(r'<a[^>]+href=["\x27]([^"\x27]+)["\x27][^>]*>(.*?)</a>',body,re.I|re.S):
  label=unescape(re.sub('<[^>]+>','',label)).strip();u=urllib.parse.urljoin(url,unescape(href))
  if not re.search('上一篇|下一篇|上一条|下一条|上一页|下一页|信息公开|办事',label) and re.search(r'附件|指南|管理办法|管理制度|申请书|下载|\.pdf|\.doc|DownloadAttach',label+' '+href,re.I):links.append({'url':u,'name':label})
 body=re.sub(r'<br\s*/?>|</(?:p|div|li|tr|h[1-6])>','\n',body,flags=re.I)
 t=unescape(re.sub('<[^>]+>',' ',body));t='\n'.join(re.sub(r'[ \t]+',' ',x).strip() for x in t.splitlines() if x.strip())
 return t,links
def get(url,referer=None):
 url=urllib.parse.quote(url,safe=':/%?=&+@')
 req=urllib.request.Request(url,headers={'User-Agent':'Mozilla/5.0','Referer':referer or url})
 try:r=urllib.request.urlopen(req,timeout=18)
 except ssl.SSLCertVerificationError:r=urllib.request.urlopen(req,context=ssl._create_unverified_context(),timeout=18)
 except urllib.error.URLError as e:
  if 'CERTIFICATE_VERIFY' in str(e):r=urllib.request.urlopen(req,context=ssl._create_unverified_context(),timeout=18)
  else:raise
 with r:b=r.read();typ=r.headers.get('Content-Type','');actual=r.url
 if b[:4]==b'%PDF':
  return '\n'.join(p.extract_text() or '' for p in PdfReader(io.BytesIO(b)).pages),[],'PDF全文'
 if b[:2]==b'PK':
  try:
   doc=Document(io.BytesIO(b));t='\n'.join(p.text for p in doc.paragraphs)+'\n'+'\n'.join(' | '.join(c.text for c in row.cells) for tab in doc.tables for row in tab.rows);return t,[],'DOCX全文'
  except:return '',[],'附件格式未解析'
 if b[:4]==bytes.fromhex('d0cf11e0'):return '',[],'旧DOC需原生解析'
 m=re.search(br'charset[=\s"\x27]+([a-zA-Z0-9_-]+)',b[:5000]);enc=m.group(1).decode() if m else 'utf8'
 try:body=b.decode(enc)
 except:body=b.decode('gb18030',errors='replace')
 t,links=htmlparse(body,actual)
 if len(t)<2000 and re.search('验证码|访问验证|无权访问|Access Denied|CAPTCHA|403 Forbidden|404 Not Found',t,re.I):return '',links,'访问受限/验证码/未找到'
 return t,links,'网页正文'
AI=json.loads((OLD/'audit_all/attachment_index.json').read_text(encoding='utf8'))
def one(ident):
 f=H/'sources'/f'{ident}.json'
 if f.exists():return json.loads(f.read_text(encoding='utf8'))
 previous=P/'sources'/f'{ident}.json'
 if previous.exists():b=json.loads(previous.read_text(encoding='utf8'))
 else:
  old=json.loads((OLD/'official'/f'{ident}.json').read_text(encoding='utf8'));b={'id':ident,'url':N[ident]['url'],'sources':[],'attempts':[],'links':[]}
  if old.get('text'):b['sources'].append({'url':old['url'],'kind':'官方正文存档','text':old['text'],'live':False})
  for a in AI:
   if ident in a.get('ids',[]) and a.get('text'):b['sources'].append({'url':a['url'],'kind':a['name'],'text':a['text'],'live':False})
 b['sources']=[s for s in b['sources'] if not re.search('上一篇|下一篇|信息公开|办事指南',s['kind'])]
 # Same-day already retrieved originals remain evidence. Retry failures and formerly excluded originals.
 refresh=not any(s.get('live') and s['url']==b['url'] for s in b['sources']) or ident in ['M035','M043','M068','M106','M114']
 if refresh:
  try:
   t,links,status=get(b['url']);b['attempts'].append({'url':b['url'],'status':status,'round':'最终收尾'})
   if t and len(t)>100:b['sources'].insert(0,{'url':b['url'],'kind':'最终收尾官方正文','text':t,'live':True});b['links']=links
  except Exception as e:b['attempts'].append({'url':b['url'],'status':str(e)[:150],'round':'最终收尾'})
 for link in b.get('links',[]):
  if not re.search('指南|管理办法|管理制度|条例|通知|年度',link['name']) or re.search('上一篇|下一篇|信息公开|办事指南',link['name']):continue
  if any(s['url']==link['url'] and s['text'] for s in b['sources']):continue
  try:
   t,_,status=get(link['url'],b['url']);b['attempts'].append({'url':link['url'],'name':link['name'],'status':status,'round':'最终收尾'})
   if t and len(t)>70:b['sources'].append({'url':link['url'],'kind':link['name'],'text':t,'live':True})
  except Exception as e:b['attempts'].append({'url':link['url'],'name':link['name'],'status':str(e)[:150],'round':'最终收尾'})
 f.write_text(json.dumps(b,ensure_ascii=False,indent=2),encoding='utf8');return b
if __name__=='__main__':
 ids=sys.argv[1].split(',') if len(sys.argv)>1 else list(N)
 with concurrent.futures.ThreadPoolExecutor(max_workers=8) as pool:
  for b in pool.map(one,ids):print(b['id'],len(b['sources']),b['attempts'][-1].get('status','')[:55] if b['attempts'] else '存档',flush=True)
 print('SOURCE_RECORDS',len(list((H/'sources').glob('*.json'))),flush=True)
