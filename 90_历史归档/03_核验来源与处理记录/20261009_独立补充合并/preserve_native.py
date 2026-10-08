from pathlib import Path
import json,zipfile,copy,xml.etree.ElementTree as E,re,hashlib
import openpyxl
H=Path(__file__).resolve().parent; R=H.parents[2]
S='http://schemas.openxmlformats.org/spreadsheetml/2006/main'; Q='{'+S+'}'; P='http://schemas.openxmlformats.org/package/2006/relationships'; O='http://schemas.openxmlformats.org/officeDocument/2006/relationships'
E.register_namespace('',S);E.register_namespace('r',O)
p=json.loads((H/'patch.json').read_text(encoding='utf8'))
with zipfile.ZipFile(H/'user_input.xlsx') as z: files={n:z.read(n) for n in z.namelist()}
with zipfile.ZipFile(H/'authored.xlsx') as z: authored={n:z.read(n) for n in z.namelist()}
wb=E.fromstring(files['xl/workbook.xml']); names=[x.get('name') for x in wb.find(Q+'sheets')]
sst=E.fromstring(authored['xl/sharedStrings.xml']) if 'xl/sharedStrings.xml' in authored else None
trees={};atrees={}
for c in p['cells']:
 idx=names.index(c['sheet'])+1;fn=f'xl/worksheets/sheet{idx}.xml'
 if fn not in trees:trees[fn]=E.fromstring(files[fn]);atrees[fn]=E.fromstring(authored[fn])
 root=trees[fn];sd=root.find(Q+'sheetData');rn=int(re.search(r'\d+',c['cell'])[0]); row=next((x for x in sd if int(x.get('r'))==rn),None)
 if row is None:
  row=copy.deepcopy(sd[-1]);row.set('r',str(rn));
  for x in list(row):row.remove(x)
  sd.append(row)
 ac=next(x for x in atrees[fn].iter(Q+'c') if x.get('r')==c['cell']); nc=copy.deepcopy(ac)
 if nc.get('t')=='s':
  si=sst[int(nc.find(Q+'v').text)]; txt=''.join(x.text or '' for x in si.iter(Q+'t'));nc.remove(nc.find(Q+'v'));nc.set('t','inlineStr');t=E.SubElement(E.SubElement(nc,Q+'is'),Q+'t');t.text=txt
 old=next((x for x in row if x.get('r')==c['cell']),None)
 if old is not None:nc.set('s',old.get('s','0'));pos=list(row).index(old);row.remove(old);row.insert(pos,nc)
 else:
  col=re.match('[A-Z]+',c['cell'])[0];template=next((x for x in sd[-2] if re.match('[A-Z]+',x.get('r'))[0]==col),None)
  if template is not None:nc.set('s',template.get('s','0'))
  row.append(nc)
 # A modified literal source URL must retain the old hyperlink only when it is still correct.
 if c['sheet']=='02项目完整详情' and c['cell'].startswith('T'):
  hp=root.find(Q+'hyperlinks'); link=next((x for x in hp if x.get('ref')==c['cell']),None) if hp is not None else None
  if link is not None:
   relfn=f'xl/worksheets/_rels/sheet{idx}.xml.rels';rels=E.fromstring(files[relfn]);rid=link.get('{'+O+'}id')
   for rel in rels:
    if rel.get('Id')==rid:rel.set('Target',c['value'])
   files[relfn]=E.tostring(rels,encoding='utf-8',xml_declaration=True)
# Remove the now verified, expired R011 from pending, shifting row-bound native features.
fn='xl/worksheets/sheet7.xml';root=E.fromstring(files[fn]);sd=root.find(Q+'sheetData');removed=p['pending_remove_row']
for row in list(sd):
 rn=int(row.get('r'))
 if rn==removed:sd.remove(row)
 elif rn>removed:
  row.set('r',str(rn-1))
  for c in row:c.set('r',re.match('[A-Z]+',c.get('r'))[0]+str(rn-1))
hp=root.find(Q+'hyperlinks')
project_link=next((copy.deepcopy(x) for x in hp if x.get('ref')=='C'+str(removed)),None)
# The retained R008 pending-view link must point to the official notices column.
rf='xl/worksheets/_rels/sheet7.xml.rels';rels=E.fromstring(files[rf])
for link in hp:
 if link.get('ref')=='F21':
  rid=link.get('{'+O+'}id')
  for rel in rels:
   if rel.get('Id')==rid:rel.set('Target','https://slmt.cqu.edu.cn/index/tzgg.htm')
files[rf]=E.tostring(rels,encoding='utf-8',xml_declaration=True)
for link in list(hp):
 rn=int(re.search(r'\d+',link.get('ref'))[0])
 if rn==removed:hp.remove(link)
 elif rn>removed:link.set('ref',re.match('[A-Z]+',link.get('ref'))[0]+str(rn-1))

if root.find(Q+'dimension') is not None:root.find(Q+'dimension').set('ref','A1:H52')
trees[fn]=root
# Historical row and audit rows use adequate height; add a real notice hyperlink.
hist=trees['xl/worksheets/sheet6.xml'];new=hist.find(Q+'sheetData')[-1];new.set('ht','96');new.set('customHeight','1')
hp=hist.find(Q+'hyperlinks');link=E.SubElement(hp,Q+'hyperlink',{'ref':'F97','{'+O+'}id':'rIdMergeR011'})
if project_link is not None:
 project_link.set('ref','C97');hp.append(project_link)
rf='xl/worksheets/_rels/sheet6.xml.rels';rels=E.fromstring(files[rf]);E.SubElement(rels,'{'+P+'}Relationship',{'Id':'rIdMergeR011','Type':O+'/hyperlink','Target':'https://keyanchu.syuct.edu.cn/info/1028/2075.htm','TargetMode':'External'});files[rf]=E.tostring(rels,encoding='utf-8',xml_declaration=True)
log=trees['xl/worksheets/sheet4.xml'];
if log.find(Q+'dimension') is not None:log.find(Q+'dimension').set('ref','A1:H'+str(p['log_end']))
for row in log.find(Q+'sheetData'):
 if int(row.get('r'))>=2386:row.set('ht','110');row.set('customHeight','1')
for fn,tree in trees.items():files[fn]=E.tostring(tree,encoding='utf-8',xml_declaration=True)
for table,ref in [(4,'A1:H'+str(p['log_end'])),(6,'A1:F97'),(7,'A1:H52')]:
 fn=f'xl/tables/table{table}.xml';root=E.fromstring(files[fn]);root.set('ref',ref)
 af=root.find(Q+'autoFilter')
 if af is not None:af.set('ref',ref)
 files[fn]=E.tostring(root,encoding='utf-8',xml_declaration=True)
out=H/'全国机器人开放课题_个人申报台账_独立复核合并版_20261009.xlsx'
with zipfile.ZipFile(out,'w',zipfile.ZIP_DEFLATED) as z:
 for fn,data in files.items():z.writestr(fn,data)
w=openpyxl.load_workbook(out);before=openpyxl.load_workbook(H/'user_input.xlsx')
for name in ['01申报总览','03机构检索覆盖','05剔除记录','08其他资助','90筛选前原始记录','91输入版本留档']:
 assert list(w[name].values)==list(before[name].values),name
ids=[x[0] for x in list(w['02项目完整详情'].values)[1:]];assert len(ids)==166 and len(set(ids))==166
assert w['01申报总览'].max_column==6 and w['01申报总览'].max_row==16
assert w['06历史项目'].max_row==97 and w['07待核项目'].max_row==52
assert all(w[n].sheet_state==before[n].sheet_state for n in w.sheetnames)
assert sum(len(s.tables) for s in w)==10
validation={'project_count':166,'current':15,'historical':96,'pending':51,'other':4,'unchanged_sheets_checked':6,'tables':10,'sha256':hashlib.sha256(out.read_bytes()).hexdigest(),'output':out.name}
(H/'validation.json').write_text(json.dumps(validation,ensure_ascii=False,indent=2),encoding='utf8');print(json.dumps(validation,ensure_ascii=False))
