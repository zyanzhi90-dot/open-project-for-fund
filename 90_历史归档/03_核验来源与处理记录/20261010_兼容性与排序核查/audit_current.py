from pathlib import Path
from openpyxl import load_workbook
from openpyxl.utils.datetime import to_excel
import json,datetime,re,hashlib
H=Path(__file__).parent;R=H.parents[2];p=R/'全国机器人开放课题_个人申报台账_当前申报清理版_20261010.xlsx'
w=load_workbook(p);native=json.loads((H/'native_current_values.json').read_text(encoding='utf-8-sig'));total=0;links=0
for n in native:
 s=w[n['sheet']];assert (s.max_row,s.max_column)==(n['rows'],n['cols'])
 expected=[]
 for row in s:
  for c in row:
   expected.append(to_excel(c.value) if isinstance(c.value,(datetime.datetime,datetime.date)) else c.value)
 assert expected==n['values'],(s.title,'Native repaired values differ')
 original_links={c.coordinate:c.hyperlink for row in s for c in row if c.hyperlink}
 assert len(original_links)==len(n['links'])
 for x in n['links']:
  old=original_links[x['ref']];assert (old.target or '')==(x['address'] or '') and (old.location or '')==(x['location'] or ''),(s.title,x['ref'])
 total+=len(expected);links+=len(original_links)
detail=list(w['02项目完整详情'].values);identifiers=[r[0] for r in detail[1:]];assert len(identifiers)==len(set(identifiers))==397
def rank(v):
 text=str(v[24] or '')+' '+str(v[2] or '')
 if 'P0 南京' in text or any(x in text for x in ['南京','滁州']):return 0
 for k,words in enumerate([['澳门'],['香港'],['广州','深圳','珠海','佛山','东莞','中山','惠州','珠三角'],['北京'],['西安'],['重庆']],1):
  if any(x in text for x in words):return k
 return 7
listed=[];sort_report=[]
for name in ['01申报总览','03历史项目','04其他资助']:
 s=w[name];keys=[]
 for i in range(2,s.max_row+1):
  rr=int(re.search(r'A(\d+)',s.cell(i,3).hyperlink.location)[1]);v=detail[rr-1];dt=s.cell(i,2).value
  assert s.cell(i,3).value==v[2]+'·'+v[3],(name,i,'Detail identity mismatch')
  assert dt==v[6],(name,i,'Deadline mismatch')
  assert s.cell(i,6).hyperlink.target==v[19],(name,i,'Notice mismatch')
  keys.append((rank(v),dt.isoformat() if isinstance(dt,datetime.datetime) else '9999'));listed.append(v[0])
  if name!='03历史项目':
   if isinstance(dt,datetime.datetime):assert dt.date()>datetime.date(2026,10,10)
   sort_report.append({'sheet':name,'row':i,'id':v[0],'rank':rank(v),'date':str(dt),'unit':s.cell(i,3).value})
 assert keys==sorted(keys),(name,'Order incorrect')
assert len(listed)==len(set(listed))==317
assert w['01申报总览'].max_column==6 and w['01申报总览'].max_row==20
assert w['04其他资助'].max_row==3
report={'file':str(p),'sha256':hashlib.sha256(p.read_bytes()).hexdigest(),'workbook_modified':False,'native_excel_default_repair_preserves_all_cell_values_and_hyperlinks':True,'checked_cells':total,'checked_hyperlinks':links,'details_unique':397,'overview_count':19,'other_current_count':2,'history_count':296,'detail_only_count':80,'date_id_notice_mapping_and_sort_verified':True,'repair_log':'Excel reports only xl/tables/table8.xml; sheet cells remain identical after repair','scope':'Existing workbook data consistency and post-repair preservation; no new web searches or claim of full notice re-verification','current_order':sort_report}
(H/'current_content_audit.json').write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf8')
print(json.dumps({k:v for k,v in report.items() if k!='current_order'},ensure_ascii=False));w.close()
