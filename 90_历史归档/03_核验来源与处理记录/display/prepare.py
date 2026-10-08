import pathlib,json,re,collections,hashlib
import openpyxl
R=pathlib.Path(__file__).parent
W=R.parents[1]
src=next((W/'outputs/20261008_robot_fund_audit').glob('*.xlsx'))
w=openpyxl.load_workbook(src,data_only=False)
cached=openpyxl.load_workbook(src,data_only=True)
home=w.worksheets[0];detail=w.worksheets[1]
norm={p['id']:p for p in json.loads((W/'tmp/rebuild/normalized.json').read_text(encoding='utf8'))['projects']}
detailmap={str(row[0].value):i for i,row in enumerate(list(detail.rows)[6:],7)}
source_rows=[list(row) for row in home.values][6:]
regions=['P0 南京及周边','P0 澳门','P0 香港','P0 广州及周边','P1 北京及周边','P1 西安及周边','P1 重庆']
def orgkey(s):
 s=str(s or '').strip()
 return re.sub(r'^(?:安徽|陕西|四川|北京|广东|江苏|浙江|重庆)[·：]','',s)
familiar=['南京信息工程大学','澳门大学','香港理工大学','华南理工大学','中国科学院自动化研究所','西北工业大学','重庆大学']
def status(s):
 s=str(s or '')
 if '仅清华校内' in s:return '限校内人员'
 if '已截止' in s:return '已截止'
 if '今日截止' in s:return '截止待确认'
 if '日期未过' in s:return '日期未过·待确认'
 if '滚动' in s:return '滚动机制·待确认'
 if '历史公告' in s:return '当期批次待核'
 if '机制或非申报公告' in s:return '当期公告待核'
 return '批次待核'
def topic(s):
 s=str(s or '').strip().replace('；','、').replace(';','、')
 s=re.sub(r'^方向/边界：','',s)
 s=re.sub(r'^与申请书的匹配方向：','',s)
 s=re.sub(r'、项目对象并非超声机器人$','',s)
 s=re.sub(r'、(?:与超声机器人直接匹配低|与超声接触控制匹配低|与机器人控制弱相关|需核公告附件资格.*)$','',s)
 s=re.sub(r'[、，](?:与(?:超声|机器人)|项目对象|必须以|核心.*主题|仅远距离|偏临床|难以直接).*$','',s)
 s=s.replace('高自由度机器人遥操作、柔性织物双臂操作','机器人遥操作、柔性操作')
 if s.startswith('先进测控/具身，'):s='先进测控、具身智能'
 if len(s)>40:s='、'.join(s.split('、')[:3])
 return s
money_overrides={
 'M142':'待核','M026':'待核','M184':'评审确定','M132':'评审确定',
 'M153':'重点≤5万元；面上≤3万元','M032':'重点45–65万元；一般15–25万元',
 'M034':'部分资助8–10万元','D-GZU-001':'重点6万元；一般3万元',
 'M051':'评审确定','M021':'目标导向≤4万元；自由探索≤2万元',
 'M040':'拟2–5万元','D-CAS-001':'一般≤10万元/项',
 'M058':'重点≥3万元；一般≥2万元','M053':'重点5万元；一般2万元',
 'M022':'重点20–30万元；面上10–12万元；青年5–8万元',
 'M062':'一般1–2万元（可浮动）','M036':'2–4万元',
 'M178':'深圳/上海分室≤10万元；一般10–20万元；重点≥50万元',
 'M014':'≤20万元/项','F002':'一般总支持10–50万元（经费5–25万元）；重点总支持50–200万元（经费25–100万元）',
 'M012':'重点预算≥20万元；一般预算<20万元','M103':'按课题类型确定',
 'M155':'上海/深圳分室≤10万元；中兴一般10–20万元、重点≥50万元',
 'M020':'一般10万元/项；企业≤100万元/项',
 'M096':'具体课题上限5、10或20万元','M095':'具体课题上限10、15、20、30或50万元',
 'F001':'总支持≥10万元（经费及软硬件）；经费≥总支持50%',
 'M039':'重点≤5万元；一般≤3万元；青年≤2万元','M030':'重点10万元；一般5万元'
}
def money(id,s):
 if id in money_overrides:return money_overrides[id]
 s=str(s or '').strip()
 if re.search('未确认|未取得|未在已取得|当期.*未|待核',s):return '待核'
 if re.search('未公布|未注明|未公开|金额.*评审',s):return '评审确定' if '评审' in s else '待核'
 s=re.split('；旧表|；资助期|；执行期|；研究期|；拟[0-9一二三四五六七八九十]+项',s)[0]
 s=re.sub(r'本年度每项|本年度|本年|每项资助|资助金额为|资助强度为','',s)
 s=re.sub(r'（[^）]*(?:校外|校内|期|年|结题|成果|署名|支付)[^）]*）','',s)
 s=re.sub(r'（不是科研匹配等级）','',s)
 s=re.sub(r'[,，；]\s*\d+(?:[—–~-]\d+)?年','',s)
 return s
records=[]
for row in source_rows:
 id=str(row[10]);p=norm[id];o=orgkey(p['org']);region=p['region']
 if o=='哈尔滨理工大学':region='P2 黑龙江·哈尔滨'
 state=status(row[2]);bucket=2 if state=='已截止' else 0 if state.startswith(('日期未过','截止待确认')) else 1
 records.append({'id':id,'org_key':o,'region':region,'familiar':o in familiar,'state_bucket':bucket,'detail_row':detailmap[id],'row':[state,row[4],str(row[3]).replace(' — ','·'),topic(row[5]),money(id,row[6]),'公告'],'original':row})
# Every institution is assigned to one regional block even if its source records differ.
orgregions={}
for x in records:
 key=x['org_key'];reg=x['region']
 rank=(regions.index(reg) if reg in regions else 7 if '长三角' in reg else 8 if '粤港澳' in reg else 9,reg)
 if key not in orgregions or rank<orgregions[key][0]:orgregions[key]=(rank,reg)
def sortkey(x):
 rank,reg=orgregions[x['org_key']]
 familiar_order=familiar.index(x['org_key']) if x['familiar'] else 99
 return (*rank,familiar_order,x['org_key'],x['state_bucket'],str(x['row'][1]),x['id'])
records.sort(key=sortkey)
tops=[]
for i,row in enumerate(list(home.values)[:6],1):
 vals=[v for v in row if v is not None]
 if vals:tops.append({'row':i,'values':list(row),'cached':list(cached.worksheets[0].values)[i-1]})
data={'source':str(src),'source_sha256':hashlib.sha256(src.read_bytes()).hexdigest(),'sheetnames':w.sheetnames,'detail_name':detail.title,'header':['申报状态','截止日期','依托单位·实验室','匹配方向','资助金额','公告'],'records':records,'top_archive':tops}
(R/'data.json').write_text(json.dumps(data,ensure_ascii=False,indent=2),encoding='utf8')
print(json.dumps({'records':len(records),'states':dict(collections.Counter(x['row'][0] for x in records)),'first_ids':[x['id'] for x in records[:12]],'max_topic':max(len(x['row'][3]) for x in records),'max_amount':max(len(x['row'][4]) for x in records)},ensure_ascii=False))
