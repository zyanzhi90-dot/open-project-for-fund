from pathlib import Path
import zipfile,re,copy
from lxml import etree
from openpyxl import load_workbook
root=Path('.'); src=root/'全国机器人开放课题_个人申报台账_申报合并与方向排序版_20261010.xlsx'
tmpdir=root/'90_历史归档'/'03_核验来源与处理记录'/'20261010_细节调整'; out=tmpdir/'全国机器人开放课题_个人申报台账_申报合并与方向排序版_20261010_细节调整.xlsx'
wb=load_workbook(src,data_only=False,read_only=True)
ws2=wb['02已截止']; ws3=wb['03所有项目完整详情']
rows2_2026=[]; rows2_other=[]
for r in range(2,ws2.max_row+1):
 d=ws2.cell(r,2).value; (rows2_2026 if getattr(d,'year',None)==2026 else rows2_other).append(r)
order2=[1]+rows2_2026+rows2_other
target_rows3=[]; other_rows3=[]
for r in range(2,ws3.max_row+1):
 lab=str(ws3.cell(r,4).value or ''); (target_rows3 if '航空航天结构力学及控制全国重点实验室' in lab else other_rows3).append(r)
order3=[1]+other_rows3+target_rows3
wb.close()
NS='http://schemas.openxmlformats.org/spreadsheetml/2006/main'
def row_reorder(xml_bytes,order):
 rootx=etree.fromstring(xml_bytes,etree.XMLParser(remove_blank_text=False)); sd=rootx.find('{%s}sheetData'%NS)
 rows={int(x.get('r')):x for x in sd.findall('{%s}row'%NS)}; used=set(order)
 ordered=[rows[r] for r in order if r in rows]+[x for r,x in sorted(rows.items()) if r not in used]
 old_to_new={old:new for new,old in enumerate(order,1)}
 for x in list(sd): sd.remove(x)
 for new_r,x in enumerate(ordered,1):
  x.set('r',str(new_r))
  for c in x.findall('.//{%s}c'%NS):
   ref=c.get('r');
   if ref: c.set('r',re.sub(r'^[A-Z]+\d+',lambda m: re.match(r'^[A-Z]+',m.group(0)).group(0)+str(new_r),ref))
  sd.append(x)
 # Hyperlinks are outside sheetData and must follow moved rows.
 for h in rootx.findall('.//{%s}hyperlink'%NS):
  ref=h.get('ref'); m=re.fullmatch(r'([A-Z]+)(\d+)',ref or '')
  if m and int(m.group(2)) in old_to_new: h.set('ref',m.group(1)+str(old_to_new[int(m.group(2))]))
 return etree.tostring(rootx,xml_declaration=True,encoding='UTF-8',standalone=True),old_to_new
def col_num(s):
 n=0
 for ch in s: n=n*26+ord(ch)-64
 return n
def col_name(n):
 s=''
 while n: n,rem=divmod(n-1,26); s=chr(65+rem)+s
 return s
def shift_ref(ref):
 def one(m):
  n=col_num(m.group(1)); row=m.group(2); n=3 if n==4 else (n-1 if n>4 else n); return f'{col_name(n)}{row}'
 return ':'.join(one(re.fullmatch(r'([A-Z]+)(\d+)',p)) if re.fullmatch(r'([A-Z]+)(\d+)',p) else p for p in ref.split(':'))
def delete_col_d(xml_bytes):
 rootx=etree.fromstring(xml_bytes,etree.XMLParser(remove_blank_text=False))
 for row in rootx.findall('.//{%s}row'%NS):
  for c in list(row):
   if c.tag!='{%s}c'%NS: continue
   m=re.fullmatch(r'([A-Z]+)(\d+)',c.get('r',''))
   if not m: continue
   n=col_num(m.group(1))
   if n==4: row.remove(c)
   elif n>4: c.set('r',f'{col_name(n-1)}{m.group(2)}')
 cols=rootx.find('{%s}cols'%NS)
 if cols is not None:
  for col in list(cols):
   mn=int(col.get('min','1')); mx=int(col.get('max',str(mn)))
   if mn==mx==4: cols.remove(col)
   elif mn>=5: col.set('min',str(mn-1)); col.set('max',str(mx-1))
   elif mn<4<=mx:
    left=copy.deepcopy(col); left.set('max','3'); right=copy.deepcopy(col); right.set('min','4'); right.set('max',str(mx-1)); i=list(cols).index(col); cols.remove(col); cols.insert(i,left); cols.insert(i+1,right)
 for el in rootx.iter():
  for attr in ('ref','sqref'):
   if attr in el.attrib: el.set(attr,' '.join(shift_ref(x) for x in el.get(attr).split()))
 return etree.tostring(rootx,xml_declaration=True,encoding='UTF-8',standalone=True)
with zipfile.ZipFile(src,'r') as zin, zipfile.ZipFile(out,'w',zipfile.ZIP_DEFLATED) as zout:
 for item in zin.infolist():
  data=zin.read(item.filename)
  if item.filename=='xl/worksheets/sheet2.xml': data,map2=row_reorder(data,order2)
  elif item.filename=='xl/worksheets/sheet3.xml': data,map3=row_reorder(data,order3)
  elif item.filename=='xl/worksheets/sheet4.xml': data=delete_col_d(data)
  elif item.filename=='xl/tables/table4.xml':
   tx=etree.fromstring(data); tx.set('ref','A6:H2536'); tx.find('{%s}autoFilter'%NS).set('ref','A6:H2536'); tcols=tx.find('{%s}tableColumns'%NS)
   for tc in list(tcols):
    if tc.get('id')=='4': tcols.remove(tc)
   for tc in tcols:
    i=int(tc.get('id')); tc.set('id',str(i-1) if i>4 else str(i))
   tcols.set('count',str(len(tcols))); data=etree.tostring(tx,xml_declaration=True,encoding='UTF-8',standalone=True)
  zout.writestr(item,data)
# Patch internal 02 hyperlinks after both maps are known.
with zipfile.ZipFile(out,'r') as zin:
 files={n:zin.read(n) for n in zin.namelist()}
rootx=etree.fromstring(files['xl/worksheets/sheet2.xml'],etree.XMLParser(remove_blank_text=False))
for h in rootx.findall('.//{%s}hyperlink'%NS):
 loc=h.get('location','')
 m=re.search(r"!A(\d+)$",loc)
 if m:
  old=int(m.group(1)); new=map3.get(old,old); h.set('location',"'03所有项目完整详情'!A"+str(new))
files['xl/worksheets/sheet2.xml']=etree.tostring(rootx,xml_declaration=True,encoding='UTF-8',standalone=True)
with zipfile.ZipFile(out,'w',zipfile.ZIP_DEFLATED) as zout:
 for n,d in files.items(): zout.writestr(n,d)
print('created',out); print('02_2026_first_count',len(rows2_2026)); print('03_moved_to_end',target_rows3)
